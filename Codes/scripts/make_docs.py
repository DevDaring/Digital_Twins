#!/usr/bin/env python3
"""Generate the submission documents from backend/reports/*.json.

Writes (every number is read from the reports or the artefact manifest at run time; nothing
is typed by hand):

* docs/figures/*.png                 charts and diagrams (matplotlib)
* docs/figures/screens/*.png         app screenshots (copied and cropped from --shots, then reused)
* docs/architecture.pptx (+ .pdf)    4 slides: system, forecast data flow, agent contract,
                                     evaluation protocol
* docs/presentation.pptx (+ .pdf)    12 slides: problem, why CGM twins miss India, idea, live demo,
                                     twin engine, sensor ladder, next-best-prick, trust and
                                     fairness, voice, safety and limits, roadmap, ask
* docs/VIDEO_SCRIPT.md               22-minute timestamped demo script (template below)
* docs/JURY_QA.md                    prepared jury answers (template below)
* ../README.md (repo root)           only the blocks between <!-- DOCS:<NAME>:START/END --> markers;
                                     the results tables between <!-- RESULTS:START/END --> belong to
                                     scripts/build_readme.py

Usage (from the Codes/ folder)::

    backend/.venv/bin/python scripts/make_docs.py                # everything, including PDFs
    backend/.venv/bin/python scripts/make_docs.py --no-pdf       # skip LibreOffice
    backend/.venv/bin/python scripts/make_docs.py --shots DIR    # refresh screenshots from DIR

Needs python-pptx, matplotlib and Pillow (all in the backend venv) and, for PDFs,
LibreOffice (``soffice``) on PATH.
"""

from __future__ import annotations

import argparse
import csv
import json
import math
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path
from typing import Any

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
from matplotlib.patches import FancyBboxPatch, Patch  # noqa: E402
from PIL import Image  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.dml.color import RGBColor  # noqa: E402
from pptx.enum.shapes import MSO_SHAPE  # noqa: E402
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN  # noqa: E402
from pptx.oxml.ns import qn  # noqa: E402
from pptx.oxml.xmlchemy import OxmlElement  # noqa: E402
from pptx.util import Emu, Inches, Pt  # noqa: E402

ROOT = Path(__file__).resolve().parents[1]
REPORTS = ROOT / "backend" / "reports"
ARTIFACTS = ROOT / "backend" / "artifacts"
FOOD = ROOT / "backend" / "data" / "food"
DOCS = ROOT / "docs"
FIG = DOCS / "figures"
SCREENS = FIG / "screens"
README = ROOT.parent / "README.md"  # repo root, outside Codes/
HOSTED = "https://pratifalan.duckdns.org"
TEAM = "Pratifalan"
MEMBER = "Koushik Deb (solo)"
INSTITUTION = "Indian Institute of Information Technology Kalyani"
CHALLENGE = "Happiest Health Digital Twin Challenge 2026"

# ----------------------------------------------------------------------------- theme
INDIGO = "#0E1230"
INDIGO_2 = "#1A1F4A"
PAPER = "#FAF6EF"
CARD = "#FFFDF8"
MARIGOLD = "#F2A33A"
TEAL = "#2BB3A3"
CORAL = "#E5604D"
VIOLET = "#7C6CF0"
INK = "#1B1F3B"
MUTED = "#5E6280"
RULE = "#E4DCCD"
# Chart series on the paper surface: marigold / teal / violet pass the dataviz validator
# (lightness band, chroma floor, CVD and normal-vision separation, contrast >= 3:1 on #FAF6EF).
# Coral is close to marigold, so it is used only for labelled status marks (targets, "not met").
C_MARIGOLD = "#C97A12"
C_TEAL = "#0E9A8C"
C_VIOLET = "#6A5AE0"
C_CORAL = "#D04A36"
C_GRAY = "#8C8778"

FONT = "Noto Sans"
INDIC_FONTS = {"deva": "Noto Sans Devanagari", "beng": "Noto Sans Bengali", "knda": "Noto Sans Kannada"}


def _register_fonts() -> None:
    """matplotlib's cache may not list system Noto fonts; register the files we need explicitly."""
    from matplotlib import font_manager

    for d in ("/usr/share/fonts/truetype/noto", "/usr/share/fonts/noto", "/usr/share/fonts/opentype/noto"):
        for stem in ("NotoSans", "NotoSansBengali", "NotoSansDevanagari", "NotoSansKannada"):
            for style in ("Regular", "Bold"):
                p = Path(d) / f"{stem}-{style}.ttf"
                if p.exists():
                    font_manager.fontManager.addfont(str(p))


_register_fonts()
plt.rcParams.update({
    "font.family": [FONT, "Noto Sans Devanagari", "Noto Sans Bengali", "Noto Sans Kannada", "DejaVu Sans"],
    "font.size": 13,
    "axes.edgecolor": RULE,
    "axes.labelcolor": MUTED,
    "axes.titlecolor": INK,
    "axes.titlesize": 15,
    "axes.titleweight": "bold",
    "axes.titlelocation": "left",
    "axes.titlepad": 12,
    "xtick.color": MUTED,
    "ytick.color": MUTED,
    "axes.grid": True,
    "grid.color": RULE,
    "grid.linewidth": 0.8,
    "axes.spines.top": False,
    "axes.spines.right": False,
    "figure.facecolor": PAPER,
    "axes.facecolor": PAPER,
    "savefig.facecolor": PAPER,
    "legend.frameon": False,
})


# ----------------------------------------------------------------------------- loading + formatting
def load_reports() -> dict[str, dict]:
    out = {}
    for p in sorted(REPORTS.glob("*.json")):
        out[p.stem] = json.loads(p.read_text(encoding="utf-8"))
    need = ["sensor_ladder", "baselines", "calibration_length", "nbp_value", "transfer", "subgroups",
            "error_grid", "calibration_curve", "agent_suite", "voice_roundtrip", "meal_photo", "summary"]
    missing = [n for n in need if n not in out]
    if missing:
        sys.exit(f"[make_docs] missing reports: {missing}. Run `make reproduce` first.")
    man = ARTIFACTS / "manifest.json"
    if not man.exists():
        sys.exit("[make_docs] backend/artifacts/manifest.json missing. Run `make experiments` first.")
    out["_manifest"] = json.loads(man.read_text(encoding="utf-8"))
    return out


def _n(x: Any) -> float:
    v = float(x)
    if math.isnan(v) or math.isinf(v):
        raise ValueError(x)
    return v


def f0(x: Any) -> str:
    return f"{_n(x):.0f}"


def f1(x: Any) -> str:
    return f"{_n(x):.1f}"


def f2(x: Any) -> str:
    return f"{_n(x):.2f}"


def f3(x: Any) -> str:
    return f"{_n(x):.3f}"


def sf1(x: Any) -> str:
    """Signed, 1 decimal, with a real minus sign."""
    v = _n(x)
    return ("+" if v > 0 else "") + f"{v:.1f}".replace("-", "−")


def sf2(x: Any) -> str:
    v = _n(x)
    return ("+" if v > 0 else "") + f"{v:.2f}".replace("-", "−")


def sf3(x: Any) -> str:
    v = _n(x)
    return ("+" if v > 0 else "") + f"{v:.3f}".replace("-", "−")


def m1(x: Any) -> str:
    """1 decimal with a real minus sign, no plus sign."""
    return f"{_n(x):.1f}".replace("-", "−")


def m2(x: Any) -> str:
    return f"{_n(x):.2f}".replace("-", "−")


def pct(x: Any) -> str:
    """Fraction -> percent; whole numbers without decimals."""
    v = 100 * _n(x)
    return f"{v:.0f}%" if abs(v - round(v)) < 1e-9 else f"{v:.1f}%"


def pct1(x: Any) -> str:
    return f"{100 * _n(x):.1f}%"


def pct0(x: Any) -> str:
    return f"{100 * _n(x):.0f}%"


def secs(ms: Any) -> str:
    return f"{_n(ms) / 1000:.1f} s"


def gen(rep: dict) -> str:
    return str(rep.get("generated_at", "?"))[:16].replace("T", " ")


class Values(dict):
    def __missing__(self, key: str) -> str:  # template typo -> loud failure
        raise KeyError(f"make_docs: unknown value '{key}'")


def hours_label(s: str) -> str:
    s = str(s)
    return s[:-2] + "+" if s.endswith("-+") else s.replace("-", "–")


def transfer_key(r: dict) -> str:
    test = "sh" if r["test"].startswith("Shanghai") else "cg"
    dom = "in" if "in-domain" in r["train"] else "tr"
    real = "_real" if "real" in r["setting"] else ""
    return f"{test}_{dom}{real}"


def agent_counts(ag: dict) -> tuple[dict[str, int], str]:
    tot: dict[str, int] = {}
    for c in ag["suite"]["counts"].values():
        for k, v in c.items():
            tot[k] = tot.get(k, 0) + v
    names = {"in_scope": "in scope", "out_of_scope": "out of scope", "unsafe": "unsafe requests",
             "emergency": "emergencies", "injection": "prompt injections"}
    return tot, ", ".join(f"{v} {names.get(k, k)}" for k, v in tot.items())


def alert_at(bins: list[dict], thr: float) -> dict[str, float]:
    """Operating point of P(>180) at a probability threshold that falls on a bin edge.

    Derived from the held-out reliability bins (count n and observed frequency per bin), so
    it is only as exact as the bins: windows inside a bin are treated alike.
    """
    tot_n = sum(b["n"] for b in bins)
    tot_pos = sum(b["n"] * b["observed"] for b in bins)
    sel = [b for b in bins if b["pred_lo"] >= thr - 1e-9]
    a_n = sum(b["n"] for b in sel)
    a_pos = sum(b["n"] * b["observed"] for b in sel)
    return {"alert_rate": a_n / tot_n, "ppv": a_pos / a_n, "sens": a_pos / tot_pos, "n": tot_n}


def values(R: dict[str, dict]) -> Values:  # noqa: C901 - one flat table of named numbers
    V = Values()
    sl = R["sensor_ladder"]
    L = {lv["level"]: lv for lv in sl["levels"]}
    tag = {"full": "full", "4": "p4", "2": "p2", "1": "p1", "0": "p0"}
    V["sl_n_patients"] = str(L["2"]["n_people"])
    V["sl_n_forecasts"] = f"{L['2']['n_forecast_windows']:,}"
    V["sl_dataset"] = sl["dataset"]
    V["validated_h"] = str(sl["validated_horizon_min"])
    for k, t in tag.items():
        lv = L[k]
        V[f"label_{t}"] = lv["label"]
        for h in ("30", "60", "90", "120"):
            V[f"rmse{h}_{t}"] = f1(lv["rmse"][h])
        V[f"rmse60_lo_{t}"], V[f"rmse60_hi_{t}"] = f1(lv["rmse60_ci"][1]), f1(lv["rmse60_ci"][2])
        V[f"cov60_{t}"] = pct1(lv["coverage90"]["60"])
        V[f"cov60_lo_{t}"], V[f"cov60_hi_{t}"] = pct1(lv["coverage90_60_ci"][1]), pct1(lv["coverage90_60_ci"][2])
        V[f"width60_{t}"] = f1(lv["width90"]["60"])
        V[f"auroc_{t}"] = f2(lv["auroc_high"])
        V[f"auroc_lo_{t}"], V[f"auroc_hi_{t}"] = f2(lv["auroc_high_ci"][1]), f2(lv["auroc_high_ci"][2])
        V[f"auprc_{t}"] = f2(lv["auprc_high"])
        V[f"brier_{t}"] = f3(lv["brier_high"])
        V[f"prev_high_{t}"] = pct1(lv["prevalence_high"])
        V[f"ece_{t}"] = f3(lv["ece_high"])
        V[f"mech_rmse60_{t}"] = f1(lv["mechanistic_only"]["rmse"]["60"])
        rc = lv["reconstruction"]
        V[f"recon_{t}"] = f1(rc["rmse"])
        V[f"clarke_ab_{t}"] = f"{rc['clarke']['A'] + rc['clarke']['B']:.1f}%"
        V[f"meal_cov60_{t}"] = pct1(lv["meal_anchored"]["coverage90"]["60"])
        V[f"meal_rmse60_{t}"] = f1(lv["meal_anchored"]["rmse"]["60"])
        V[f"meal_n_{t}"] = f"{lv['meal_anchored']['n_forecast_windows']:,}"
        V[f"t2d_rmse60_{t}"] = f1(lv["t2d_only"]["rmse"]["60"])
        V[f"t2d_n_{t}"] = str(lv["t2d_only"]["n_people"])
        V[f"abstain_{t}"] = pct1(lv["abstain_rate"])
        V[f"rmse60_abst_{t}"] = f1(lv["abstain_vs_not"]["rmse60_abstained"])
        V[f"rmse60_conf_{t}"] = f1(lv["abstain_vs_not"]["rmse60_confident"])
        ex = lv["exploratory_horizons"]["horizons"]
        for h in ("180", "240"):
            V[f"x{h}_rmse_{t}"] = f1(ex[h]["rmse"])
            V[f"x{h}_cov_{t}"] = pct1(ex[h]["coverage90"])
            V[f"x{h}_width_{t}"] = f1(ex[h]["width90"])
    le = L["2"]["low_events"]
    V["low_windows"], V["low_patients"] = str(le["positive_windows"]), str(le["patients_with_lows"])
    V["low_prev"] = pct1(le["prevalence"])
    V["d_rmse60_p2_p0"] = f1(_n(L["0"]["rmse"]["60"]) - _n(L["2"]["rmse"]["60"]))
    V["d_rmse60_p2_full"] = f1(_n(L["2"]["rmse"]["60"]) - _n(L["full"]["rmse"]["60"]))

    rp = sl["recent_prick_value"]
    V["rp_n"] = f"{rp['n_forecast_windows']:,}"
    V["rp_people"] = str(rp["n_people"])
    for h in ("30", "60", "90", "120"):
        V[f"rp_with{h}"] = f1(rp["rmse_with_prick"][h])
        V[f"rp_without{h}"] = f1(rp["rmse_without"][h])
    d = rp["rmse60_difference_ci"]
    V["rp_diff"], V["rp_lo"], V["rp_hi"] = m1(d[0]), m1(d[1]), m1(d[2])

    pb = sl["paired_band_width"]
    for sub, t in (("recent_prick", "near"), ("control_far_from_prick", "far")):
        s = pb[sub]
        V[f"pbw_{t}_n"] = f"{s['n_forecast_windows']:,}"
        V[f"pbw_{t}_subset"] = s["subset"]
        for h, r in s["by_horizon"].items():
            V[f"pbw_{t}_with{h}"] = f1(r["mean_width_with_pricks"])
            V[f"pbw_{t}_without{h}"] = f1(r["mean_width_without_glucose"])
            V[f"pbw_{t}_d{h}"] = sf1(r["paired_difference"])
            V[f"pbw_{t}_lo{h}"], V[f"pbw_{t}_hi{h}"] = m1(r["paired_difference_ci95"][0]), m1(r["paired_difference_ci95"][1])
            V[f"pbw_{t}_covw{h}"] = pct1(r["coverage90_with_pricks"])
            V[f"pbw_{t}_covo{h}"] = pct1(r["coverage90_without_glucose"])
    bins = sl["by_hours_since_reading"]
    V["bh_widths"] = " / ".join(f1(b["width90_60"]) for b in bins)
    V["bh_labels"] = ", ".join(hours_label(b["hours_since_reading"]) + " h" for b in bins)

    # baselines and paired differences
    B = {r["level"]: r for r in R["baselines"]["rows"]}
    for lvl, t in (("2", "p2"), ("0", "p0"), ("full", "full")):
        for meth, m in B[lvl]["methods"].items():
            V[f"b_{meth}_{t}"] = f1(m["rmse"]["60"])
        for meth, pdiff in B[lvl]["paired_vs_full_hybrid_60min"].items():
            V[f"bp_{meth}_{t}"] = sf2(pdiff["difference"])
            V[f"bp_{meth}_lo_{t}"], V[f"bp_{meth}_hi_{t}"] = m2(pdiff["difference_ci95"][0]), m2(pdiff["difference_ci95"][1])
            V[f"bp_{meth}_sig_{t}"] = "yes" if pdiff["ci_excludes_zero"] else "no"
    V["b_lgbm_auroc_p2"] = f2(B["2"]["methods"]["lightgbm_sparse_only"]["auroc_high"])

    for r in R["calibration_length"]["rows"]:
        d_ = r["cgm_days"]
        V[f"cl{d_}_rmse60"] = f1(r["rmse"]["60"])
        V[f"cl{d_}_cov60"] = pct1(r["coverage90"]["60"])
        V[f"cl{d_}_auroc"] = f2(r["auroc_high"])

    nb = R["nbp_value"]
    arms = nb["arms"]
    V["nbp_fixed_name"] = next(k for k in arms if k.startswith("fixed"))
    for k, t in ((V["nbp_fixed_name"], "fixed"), ("random", "random"), ("twin-recommended", "twin")):
        V[f"nbp_{t}_rmse60"] = f1(arms[k]["rmse"]["60"])
        V[f"nbp_{t}_recon"] = f1(arms[k]["reconstruction"]["rmse"])
    for k, t in (("twin_minus_fixed", "tf"), ("twin_minus_random", "tr")):
        pdiff = nb["paired_differences"][k]
        V[f"nbp_{t}_mean"] = sf2(pdiff["mean_recon_rmse_change"])
        V[f"nbp_{t}_lo"], V[f"nbp_{t}_hi"] = m2(pdiff["ci95"][0]), m2(pdiff["ci95"][1])
        V[f"nbp_{t}_improved"] = str(pdiff["patients_improved"])
        V[f"nbp_{t}_n"] = str(pdiff["n_people"])

    tr = R["transfer"]
    for r in tr["rows"]:
        key = transfer_key(r)
        V[f"{key}_rmse60"] = f1(r["rmse"]["60"])
        V[f"{key}_lo"], V[f"{key}_hi"] = f1(r["rmse60_ci"][1]), f1(r["rmse60_ci"][2])
        V[f"{key}_cov60"] = pct1(r["coverage90"]["60"])
        V[f"{key}_width60"] = f1(r["width90"]["60"])
        V[f"{key}_auroc"] = f2(r["auroc_high"])
        if r["test"] == "ShanghaiT2DM":
            V["sh_people"], V["sh_recordings"] = str(r["n_people"]), str(r["n_recordings"])
    V["gap_sh"] = f1(tr["gap_rmse60"]["shanghai_2pricks"])
    V["gap_sh_real"] = f1(tr["gap_rmse60"]["shanghai_real_pricks"])
    V["gap_cg"] = f1(tr["gap_rmse60"]["cgmacros_2pricks"])
    ins = tr["shanghai_by_insulin"]
    for k, t in (("insulin_users", "ins"), ("no_insulin", "noins")):
        V[f"{t}_people"], V[f"{t}_rec"] = str(ins[k]["n_people"]), str(ins[k]["n_recordings"])
        V[f"{t}_rmse60"] = f1(ins[k]["rmse"]["60"])
        V[f"{t}_cov60"] = pct1(ins[k]["coverage90"]["60"])

    sg = R["subgroups"]
    rows = sg["rows"]
    V["sg_overall_rmse60"] = f1(sg["overall"]["rmse60"])
    V["sg_overall_cov60"] = pct1(sg["overall"]["coverage90_60"])
    V["sg_n_groups"] = str(len(rows))
    V["sg_n_flagged"] = str(sum(1 for r in rows if r["flagged"]))
    lo = min(rows, key=lambda r: r["coverage90_60"])
    V["sg_lowcov_group"] = f"{lo['dimension']} {lo['group']}"
    V["sg_lowcov"], V["sg_lowcov_n"] = pct1(lo["coverage90_60"]), str(lo["n_people"])
    wa = min(rows, key=lambda r: r["auroc_high"])
    V["sg_lowauroc_group"] = f"{wa['dimension']} {wa['group']}"
    V["sg_lowauroc"], V["sg_lowauroc_n"] = f2(wa["auroc_high"]), str(wa["n_people"])
    wr = max(rows, key=lambda r: r["rmse60"])
    V["sg_hirmse_group"] = f"{wr['dimension']} {wr['group']}"
    V["sg_hirmse"], V["sg_hirmse_n"] = f1(wr["rmse60"]), str(wr["n_people"])
    sm = min(rows, key=lambda r: r["n_people"])
    V["sg_small_group"] = f"{sm['dimension']} {sm['group']}"
    V["sg_small_n"] = str(sm["n_people"])
    V["sg_flag_rule"] = sg["flag_rule"]

    eg = R["error_grid"]["levels"]
    for k, t in tag.items():
        f60 = eg[k]["forecast_60"]
        V[f"eg_f60_ab_{t}"] = f"{f60['A'] + f60['B']:.1f}%"
        V[f"eg_f60_a_{t}"] = f"{f60['A']:.1f}%"

    cc = R["calibration_curve"]["levels"]
    V["ece_cal_p2"], V["ece_uncal_p2"] = f3(cc["2"]["ece_high"]), f3(cc["2"]["ece_high_uncalibrated"])
    op = alert_at(cc["2"]["high"], 0.5)
    V["al_rate"], V["al_ppv"], V["al_false"], V["al_sens"] = (pct0(op["alert_rate"]), pct0(op["ppv"]),
                                                            pct0(1 - op["ppv"]), pct0(op["sens"]))
    V["al_n"] = f"{op['n']:,}"

    # protocol change (summary.json) and manifest
    pc = R["summary"]["protocol_change"]
    o, n_, dl = pc["old"], pc["new"], pc["delta_new_minus_old"]
    V["pc_rmse_old"], V["pc_rmse_new"], V["pc_rmse_d"] = f2(o["ladder_2"]["rmse60"]), f2(n_["ladder_2"]["rmse60"]), sf2(dl["ladder_2"]["rmse60"])
    V["pc_auroc_old"], V["pc_auroc_new"], V["pc_auroc_d"] = f3(o["ladder_2"]["auroc_high"]), f3(n_["ladder_2"]["auroc_high"]), sf3(dl["ladder_2"]["auroc_high"])
    V["pc_cov_old"], V["pc_cov_new"] = pct1(o["ladder_2"]["coverage90_60"]), pct1(n_["ladder_2"]["coverage90_60"])
    V["pc_cov_d"] = sf1(100 * dl["ladder_2"]["coverage90_60"]) + " pts"
    V["pc_w_old"], V["pc_w_new"], V["pc_w_d"] = f1(o["ladder_2"]["width90_60"]), f1(n_["ladder_2"]["width90_60"]), sf1(dl["ladder_2"]["width90_60"])
    V["pc_full_old"], V["pc_full_new"], V["pc_full_d"] = f2(o["full"]["rmse60"]), f2(n_["full"]["rmse60"]), sf2(dl["full"]["rmse60"])
    V["pc_gap_old"], V["pc_gap_new"], V["pc_gap_d"] = (f2(o["transfer_gap_rmse60_shanghai_2pricks"]),
                                                       f2(n_["transfer_gap_rmse60_shanghai_2pricks"]),
                                                       sf2(dl["transfer_gap_rmse60_shanghai_2pricks"]))
    man = R["_manifest"]
    V["man_checks"] = str(man["leakage_audit"]["n_checks"])
    V["man_passed"] = "passed" if man["leakage_audit"]["passed"] else "FAILED"
    V["man_version"] = man["model_version"]
    V["man_folds"] = str(man["config"]["n_folds"])
    V["prod_people"] = str(man["hybrid_trained_on"]["n_people"])
    V["prod_excluded"] = ", ".join(man["hybrid_trained_on"]["excluded_people"])
    V["n_windows"] = f"{R['summary']['n_forecast_windows']:,}"

    ag = R["agent_suite"]
    T, Lv = ag["modes"]["template"]["summary"], ag["modes"]["live"]["summary"]
    V["ag_n"] = str(ag["suite"]["n_prompts"])
    tot, V["ag_counts_cat"] = agent_counts(ag)
    for k, v in tot.items():
        V[f"ag_cnt_{k}"] = str(v)
    for pre, S in (("t", T), ("l", Lv)):
        for k in ("case_pass_rate", "grounded_or_fallback_rate", "unsafe_advice_rate", "block_rate_unsafe",
                  "emergency_recall", "false_emergency_rate", "injection_resisted_rate",
                  "out_of_scope_handled_rate", "language_match_rate", "intent_accuracy",
                  "expected_tool_called_rate", "fallback_rate", "llm_reply_rate",
                  "frequency_phrase_ok_rate", "no_stray_english_rate", "false_block_rate_in_scope"):
            V[f"{pre}_{k}"] = pct(S[k])
        V[f"{pre}_median_ms"] = f0(S["median_latency_ms"]) if S["median_latency_ms"] >= 10 else f1(S["median_latency_ms"])
        V[f"{pre}_p90_ms"] = f0(S["p90_latency_ms"])
    ho = ag["held_out"]
    V["ho_n"] = str(ho["template"]["n"])
    for pre, mode in (("t", "template"), ("l", "live")):
        V[f"ho_{pre}_pass"] = pct(ho[mode]["case_pass_rate"])
        V[f"ho_{pre}_emerg"] = pct(ho[mode]["emergency_recall"])
        V[f"ho_{pre}_block"] = pct(ho[mode]["block_rate_unsafe"])
        V[f"ho_{pre}_unsafe"] = pct(ho[mode]["unsafe_advice_rate"])
    hb = ag["held_out_before_fix"]
    V["hb_t_emerg"], V["hb_t_block"] = pct1(hb["template"]["emergency_recall"]), pct1(hb["template"]["block_rate_unsafe"])
    V["hb_t_unsafe"] = pct(hb["template"]["unsafe_advice_rate"])
    V["hb_l_emerg"], V["hb_l_block"] = pct(hb["live"]["emergency_recall"]), pct(hb["live"]["block_rate_unsafe"])
    V["hb_n_fail"] = str(len(hb["template_failures"]))
    V["hb_fails"] = ", ".join(f"`{x}`" for x in hb["template_failures"])
    V["ag_live_model"] = ag["modes"]["live"]["models"]["planner"][0]
    V["ag_n_issues"] = str(len(ag["known_issues"]))

    vr = R["voice_roundtrip"]
    vh = vr["headline"]
    vs = vr["runs"]["live"]["summary"]
    V["v_trips"] = str(vh["n_trips"])
    V["v_target"] = f"{vr['target_ttfa_s']:.1f} s"
    V["v_ttfa"], V["v_ttfa_max"] = secs(vh["median_ttfa_ms"]), secs(vs["max_ttfa_ms"])
    V["v_streamed"] = secs(vs["median_ttfa_if_streamed_ms"])
    V["v_stt"], V["v_agent"], V["v_tts"] = secs(vh["median_stt_ms"]), secs(vh["median_agent_ms"]), secs(vh["median_tts_reply_ms"])
    V["v_share"] = pct(vh["share_meeting_2_5s"])
    V["v_cer"] = f3(vh["median_cer"])
    V["v_intent"] = pct1(vs["intent_preserved_rate"])
    for lang, s in vr["runs"]["live"]["by_language"].items():
        V[f"v_ttfa_{lang[:2]}"] = secs(s["median_ttfa_ms"])
        V[f"v_intent_{lang[:2]}"] = f"{round(s['intent_preserved_rate'] * s['n'])} of {s['n']}"
    v3 = vr["runs"].get("final_v3")
    if v3:
        s3 = v3["summary"]
        V["v3_ttfa"], V["v3_max"] = secs(s3["median_ttfa_ms"]), secs(s3["max_ttfa_ms"])
        V["v3_intent"] = pct1(s3["intent_preserved_rate"])
        k3 = v3["by_language"]["kn-IN"]
        V["v3_intent_kn"] = f"{round(k3['intent_preserved_rate'] * k3['n'])} of {k3['n']}"
        V["v3_share"] = pct(s3["share_meeting_2_5s"])

    mp = R["meal_photo"]
    mm = mp["metrics"]
    V["mp_n"], V["mp_people"] = str(mp["n_photos"]), str(mp["n_participants"])
    V["mp_model"] = next(iter(mp["models"]))
    V["mp_mae"], V["mp_mae_lo"], V["mp_mae_hi"] = f1(mm["mae_g"]["value"]), f1(mm["mae_g"]["ci95"][0]), f1(mm["mae_g"]["ci95"][1])
    V["mp_medae"] = f1(mm["median_abs_error_g"]["value"])
    V["mp_medape"] = f0(mm["median_ape_pct"]["value"]) + "%"
    V["mp_bias"] = m1(mm["bias_g"]["value"])
    V["mp_bias_lo"], V["mp_bias_hi"] = m1(mm["bias_g"]["ci95"][0]), m1(mm["bias_g"]["ci95"][1])
    V["mp_loa_lo"], V["mp_loa_hi"] = m1(mm["loa_lower_g"]["value"]), sf1(mm["loa_upper_g"]["value"])
    V["mp_rho"] = f2(mm["spearman_rho"]["value"])
    V["mp_rho_lo"], V["mp_rho_hi"] = m2(mm["spearman_rho"]["ci95"][0]), m2(mm["spearman_rho"]["ci95"][1])
    V["mp_w20"] = pct(round(mm["within_20pct"]["value"], 2))
    ex = mp["metrics_excluding_repeated_drink_logs"]
    V["mpx_n"], V["mpx_excl"] = str(ex["n_photos"]), str(ex["n_excluded"])
    V["mpx_rho"] = f2(ex["metrics"]["spearman_rho"]["value"])
    V["mpx_rho_lo"], V["mpx_rho_hi"] = f2(ex["metrics"]["spearman_rho"]["ci95"][0]), f2(ex["metrics"]["spearman_rho"]["ci95"][1])
    fi = mp["forecast_impact"]
    tab = {(r["subset"].startswith("all"), r["key"]): r for r in fi["table"]}
    r60, rm = tab[(True, "rmse_60")], tab[(False, "rmse_60")]
    V["mpf_n"] = f"{r60['n_forecasts']:,}"
    V["mpf_logged60"], V["mpf_photo60"] = f1(r60["logged_carbs"]), f1(r60["photo_carbs"])
    V["mpf_diff60"], V["mpf_lo60"], V["mpf_hi60"] = sf1(r60["difference"]), f1(r60["difference_ci95"][0]), f1(r60["difference_ci95"][1])
    V["mpf_meal_n"] = str(rm["n_forecasts"])
    V["mpf_meal_diff60"], V["mpf_meal_lo60"], V["mpf_meal_hi60"] = sf1(rm["difference"]), f1(rm["difference_ci95"][0]), f1(rm["difference_ci95"][1])
    V["mpf_mcov_logged"] = pct1(fi["meal_origins"]["coverage90"]["60"]["logged_carbs"])
    V["mpf_mcov_photo"] = pct1(fi["meal_origins"]["coverage90"]["60"]["photo_carbs"])
    ind = mp["indian_samples"]
    V["mpi_n"] = str(ind["n_photos"])
    V["mpi_recall"], V["mpi_precision"] = pct(round(ind["mean_detection_recall"], 2)), pct(round(ind["mean_detection_precision"], 2))

    # Indian food table and curated swaps
    with (FOOD / "indian_foods.csv").open(encoding="utf-8") as fh:
        V["food_n"] = str(sum(1 for _ in csv.DictReader(fh)))
    V["swaps_n"] = str(len(json.loads((FOOD / "swaps.json").read_text(encoding="utf-8"))))

    for name, rep in R.items():
        if not name.startswith("_"):
            V[f"gen_{name}"] = gen(rep)
    return V


# ----------------------------------------------------------------------------- charts
def save(fig: plt.Figure, name: str) -> Path:
    p = FIG / name
    fig.savefig(p, dpi=200, bbox_inches="tight", pad_inches=0.15)
    plt.close(fig)
    return p


def _src(fig: plt.Figure, text: str, y: float = -0.04) -> None:
    fig.text(0.0, y, text, fontsize=10.5, color=MUTED, ha="left", va="top")


def chart_ladder(R: dict) -> Path:
    sl = R["sensor_ladder"]
    L = sl["levels"]
    short = {"full": "Full\nCGM", "4": "4 pricks\n/day", "2": "2 pricks\n/day", "1": "1 prick\n/day", "0": "No\nglucose"}
    fig, axes = plt.subplots(1, 3, figsize=(14.0, 4.7), gridspec_kw={"width_ratios": [1.05, 1, 1.1], "wspace": 0.34})

    ax = axes[0]
    xs = list(range(len(L)))
    pts = [lv["rmse"]["60"] for lv in L]
    los = [lv["rmse60_ci"][1] for lv in L]
    his = [lv["rmse60_ci"][2] for lv in L]
    ax.vlines(xs, los, his, color=C_MARIGOLD, lw=2.5, alpha=0.5)
    ax.plot(xs, pts, "o", color=C_MARIGOLD, ms=10, mec=PAPER, mew=2)
    for x, p, lo_ in zip(xs, pts, los):
        ax.text(x, lo_ - 0.5, f1(p), va="top", ha="center", fontsize=12, color=INK)
    ax.set_xticks(xs, [short[lv["level"]] for lv in L], fontsize=11)
    ax.set_ylabel("RMSE at 60 min (mg/dL)")
    ax.set_title("60-min error, 95% CI")
    ax.set_ylim(min(los) - 3.5, max(his) + 1.5)
    ax.set_xlim(-0.5, len(L) - 0.3)
    ax.grid(axis="x", visible=False)

    ax = axes[1]
    rp = sl["recent_prick_value"]
    hs = ["30", "60", "90", "120"]
    hx = [int(h) for h in hs]
    w = [rp["rmse_with_prick"][h] for h in hs]
    wo = [rp["rmse_without"][h] for h in hs]
    ax.plot(hx, wo, "-o", color=C_VIOLET, lw=2, ms=8, mec=PAPER, mew=2)
    ax.plot(hx, w, "-o", color=C_TEAL, lw=2, ms=8, mec=PAPER, mew=2)
    ax.set_xticks(hx, [f"{h} min" for h in hs])
    ax.set_ylabel("RMSE (mg/dL)")
    ax.set_ylim(min(w) - 4, max(wo) + 3)
    ax.set_title(f"Value of a recent prick (n={rp['n_forecast_windows']:,})")
    ax.legend(handles=[Patch(color=C_TEAL, label="prick in last 90 min"), Patch(color=C_VIOLET, label="no glucose data")],
              loc="upper center", bbox_to_anchor=(0.5, -0.13), ncol=2, fontsize=10)
    ax.grid(axis="x", visible=False)

    ax = axes[2]
    pb = sl["paired_band_width"]
    series = [("recent_prick", C_TEAL, "≤ 90 min after a prick", -4), ("control_far_from_prick", C_VIOLET, "≥ 8 h after a prick", 4)]
    for sub, col, lab, dx in series:
        bh = pb[sub]["by_horizon"]
        xs_ = [int(h) + dx for h in hs]
        ys_ = [bh[h]["paired_difference"] for h in hs]
        lo_ = [bh[h]["paired_difference_ci95"][0] for h in hs]
        hi_ = [bh[h]["paired_difference_ci95"][1] for h in hs]
        ax.vlines(xs_, lo_, hi_, color=col, lw=2.5, alpha=0.5)
        ax.plot(xs_, ys_, "-o", color=col, lw=2, ms=8, mec=PAPER, mew=2, label=f"{lab} (n={pb[sub]['n_forecast_windows']:,})")
    ax.axhline(0, color=INK, lw=1.1)
    ax.text(124, 0.8, "same width", fontsize=10, color=MUTED, ha="right", va="bottom")
    ax.set_xticks(hx, [f"{h} min" for h in hs])
    ax.set_ylabel("Band width: 2 pricks − none (mg/dL)")
    ax.set_title("90% band, paired at the same origins")
    ax.legend(loc="upper center", bbox_to_anchor=(0.5, -0.13), ncol=1, fontsize=10)
    ax.grid(axis="x", visible=False)
    _src(fig, f"Source: backend/reports/sensor_ladder.json ({gen(sl)}). CGMacros, {L[0]['n_people']} adults, 5-day CGM calibration, "
              "nested person-wise 5-fold CV, 95% CIs by bootstrap over people; pricks simulated from the hidden CGM.", y=-0.16)
    return save(fig, "result_sensor_ladder.png")


def chart_nbp(R: dict) -> Path:
    nb = R["nbp_value"]
    pdif = nb["paired_differences"]
    labels = {"twin_minus_fixed": "Twin-recommended\nvs fixed 07:00 & 22:00", "twin_minus_random": "Twin-recommended\nvs random times"}
    fig, ax = plt.subplots(figsize=(9.5, 3.6))
    ys = list(range(len(pdif)))[::-1]
    for y, (k, d) in zip(ys, pdif.items()):
        lo, hi = d["ci95"]
        ax.hlines(y, lo, hi, color=C_VIOLET, lw=3, alpha=0.6)
        ax.plot(d["mean_recon_rmse_change"], y, "o", color=C_VIOLET, ms=11, mec=PAPER, mew=2)
        ax.text(hi + 0.04, y, f"{sf2(d['mean_recon_rmse_change'])}  ({m2(lo)} to {m2(hi)});  "
                f"{d['patients_improved']}/{d['n_people']} people improved", va="center", fontsize=12, color=INK)
    ax.axvline(0, color=INK, lw=1.2)
    ax.text(0.02, len(pdif) - 0.45, "← twin better", fontsize=10, color=MUTED, ha="right")
    ax.set_yticks(ys, [labels.get(k, k) for k in pdif], fontsize=12)
    allv = [v for d in pdif.values() for v in d["ci95"]]
    ax.set_xlim(min(allv) - 0.25, max(allv) + 1.9)
    ax.set_ylim(-0.7, len(pdif) - 0.3)
    ax.set_xlabel("Change in reconstruction RMSE (mg/dL); negative = twin better")
    ax.set_title("Same 2 pricks a day, different timing: no significant gain")
    ax.grid(axis="y", visible=False)
    _src(fig, f"Source: backend/reports/nbp_value.json ({gen(nb)}). Paired over people, 95% bootstrap CI.", y=-0.1)
    return save(fig, "result_nbp.png")


def chart_trust(R: dict) -> Path:
    plt.rcParams["font.size"] = 14
    fig, axes = plt.subplots(1, 4, figsize=(17.5, 5.0), gridspec_kw={"width_ratios": [1, 1, 1.15, 1.05], "wspace": 0.6})
    cc = R["calibration_curve"]["levels"]["2"]
    ax = axes[0]
    ax.plot([0, 1], [0, 1], color=C_GRAY, lw=1, ls="--")
    for key, col, lab in (("high_uncalibrated", C_VIOLET, "raw particles"), ("high", C_MARIGOLD, "stacked + isotonic")):
        pts_ = cc[key]
        ax.plot([p["mean_pred"] for p in pts_], [p["observed"] for p in pts_], "-o", color=col, lw=2, ms=7,
                mec=PAPER, mew=1.5, label=f"{lab}, ECE {f3(cc['ece_' + key])}")
    ax.set_xlabel("Predicted P(>180 in 2 h)")
    ax.set_ylabel("Observed frequency")
    ax.set_title("Calibration, 2 pricks/day")
    ax.legend(loc="upper left", fontsize=11)
    ax.set_xlim(0, 1)
    ax.set_ylim(0, 1)

    ax = axes[1]
    L = R["sensor_ladder"]["levels"]
    short = {"full": "Full CGM", "4": "4 pricks/day", "2": "2 pricks/day", "1": "1 prick/day", "0": "No glucose"}
    ys = list(range(len(L)))[::-1]
    for y, lv in zip(ys, L):
        wv, cv = lv["width90"]["60"], lv["coverage90"]["60"]
        ax.barh(y, wv, 0.6, color=C_TEAL, edgecolor=PAPER, linewidth=2)
        ax.text(wv - 3, y, f"{f0(wv)} mg/dL · {pct1(cv)}", va="center", ha="right", fontsize=11, color="#FFFFFF",
                fontweight="bold")
    ax.set_yticks(ys, [short[lv["level"]] for lv in L], fontsize=11)
    ax.set_xlim(0, 125)
    ax.set_xlabel("90% band width at 60 min (mg/dL)")
    ax.set_title("Width (bar) and coverage")
    ax.grid(axis="y", visible=False)

    ax = axes[2]
    rows = R["subgroups"]["rows"]
    ys = list(range(len(rows)))[::-1]
    ax.axvspan(0.70, 0.82, color=C_CORAL, alpha=0.07)
    ax.axvline(0.9, color=C_GRAY, lw=1, ls="--")
    ax.axvline(0.82, color=C_CORAL, lw=1)
    for y, r in zip(ys, rows):
        ax.plot(r["coverage90_60"], y, "o", color=C_TEAL, ms=9, mec=PAPER, mew=1.5)
        ax.text(r["coverage90_60"] + 0.006, y, f"{pct1(r['coverage90_60'])}  n={r['n_people']}", va="center",
                fontsize=10.5, color=INK)
    ax.set_yticks(ys, [f"{r['dimension'].replace(' (Asian cut-offs)', '')} {r['group'].split(' (')[0]}" for r in rows], fontsize=10.5)
    ax.set_xlim(0.78, 1.02)
    ax.xaxis.set_major_formatter(matplotlib.ticker.FuncFormatter(lambda v, _p: f"{100 * v:.0f}%"))
    ax.set_xlabel("Coverage at 60 min (flag < 82%)")
    ax.set_title("Subgroups, 2 pricks/day")
    ax.set_ylim(-0.7, len(rows) - 0.4)
    ax.grid(axis="y", visible=False)

    ax = axes[3]
    tr = {transfer_key(r): r for r in R["transfer"]["rows"]}
    sh = tr["sh_in"]
    groups = [("Shanghai\nreal", "sh_in_real", "sh_tr_real"), ("Shanghai\nsimulated", "sh_in", "sh_tr"),
              ("CGMacros\nsimulated", "cg_in", "cg_tr")]
    bw = 0.36
    for i, (_lab, a, b) in enumerate(groups):
        va, vb = tr[a]["rmse"]["60"], tr[b]["rmse"]["60"]
        ax.bar(i - bw / 2 - 0.01, va, bw, color=C_TEAL, edgecolor=PAPER, linewidth=2)
        ax.bar(i + bw / 2 + 0.01, vb, bw, color=C_VIOLET, edgecolor=PAPER, linewidth=2)
        ax.text(i - bw / 2, va + 0.8, f1(va), ha="center", fontsize=11, color=INK)
        ax.text(i + bw / 2, vb + 0.8, f1(vb), ha="center", fontsize=11, color=INK)
    ax.set_xticks(range(len(groups)), [g[0] for g in groups], fontsize=11)
    ax.set_ylabel("RMSE at 60 min (mg/dL)")
    ax.set_title("Transfer between populations")
    ax.set_ylim(0, max(r["rmse"]["60"] for r in tr.values()) * 1.32)
    ax.legend(handles=[Patch(color=C_TEAL, label="fitted in-domain"), Patch(color=C_VIOLET, label="fitted on the other dataset")],
              loc="upper right", fontsize=10.5)
    ax.text(0.5, -0.30, f"ShanghaiT2DM: {sh['n_people']} people, {sh['n_recordings']} recordings; CIs resample people",
            transform=ax.transAxes, ha="center", fontsize=10, color=MUTED)
    ax.grid(axis="x", visible=False)
    _src(fig, f"Sources: backend/reports/calibration_curve.json, sensor_ladder.json, subgroups.json, transfer.json "
              f"({gen(R['transfer'])}). Held-out predictions from the nested outer folds.", y=-0.1)
    out = save(fig, "result_trust.png")
    plt.rcParams["font.size"] = 13
    return out


def chart_voice(R: dict) -> Path:
    ag = R["agent_suite"]
    T, Lv = ag["modes"]["template"]["summary"], ag["modes"]["live"]["summary"]
    HB = ag["held_out_before_fix"]["template"]
    metrics = [("case_pass_rate", "All checks passed"), ("emergency_recall", "Emergency recall"),
               ("block_rate_unsafe", "Dose / medicine requests blocked"), ("grounded_or_fallback_rate", "Numbers grounded (or template)"),
               ("language_match_rate", "Reply in the user's language"), ("expected_tool_called_rate", "Expected twin tool called")]
    fig, axes = plt.subplots(1, 2, figsize=(13.5, 5.0), gridspec_kw={"width_ratios": [1.25, 1], "wspace": 0.55})
    ax = axes[0]
    ys = list(range(len(metrics)))[::-1]
    bh = 0.27
    for y, (k, _lab) in zip(ys, metrics):
        ax.barh(y + bh, T[k], bh * 0.92, color=C_VIOLET, edgecolor=PAPER, linewidth=1.5)
        ax.barh(y, Lv[k], bh * 0.92, color=C_MARIGOLD, edgecolor=PAPER, linewidth=1.5)
        ax.text(T[k] + 0.01, y + bh, pct(T[k]), va="center", fontsize=9.5, color=INK)
        ax.text(Lv[k] + 0.01, y, pct(Lv[k]), va="center", fontsize=9.5, color=INK)
        if k in HB:
            ax.barh(y - bh, HB[k], bh * 0.92, color=C_GRAY, edgecolor=PAPER, linewidth=1.5)
            ax.text(HB[k] + 0.01, y - bh, pct1(HB[k]) + " before fix", va="center", fontsize=9.5, color=INK)
    ax.set_yticks(ys, [m[1] for m in metrics], fontsize=11)
    ax.set_xlim(0, 1.3)
    ax.set_xticks([0, 0.25, 0.5, 0.75, 1.0], ["0", "25%", "50%", "75%", "100%"])
    ax.set_title(f"Agent suite: {ag['suite']['n_prompts']} prompts, 4 languages")
    ax.legend(handles=[Patch(color=C_VIOLET, label="offline templates"), Patch(color=C_MARIGOLD, label="live LLM"),
                       Patch(color=C_GRAY, label=f"held-out {ag['held_out_before_fix']['template']['n']}, templates, before fix")],
              loc="upper center", bbox_to_anchor=(0.4, -0.08), ncol=3, fontsize=10)
    ax.grid(axis="y", visible=False)

    ax = axes[1]
    vr = R["voice_roundtrip"]["runs"]["live"]["by_language"]
    order = [lg for lg in ("en-IN", "hi-IN", "bn-IN", "kn-IN") if lg in vr]
    names = {"en-IN": "English", "hi-IN": "Hindi", "bn-IN": "Bengali", "kn-IN": "Kannada"}  # matplotlib cannot shape Indic
    ys = list(range(len(order)))[::-1]
    for y, lg in zip(ys, order):
        s = vr[lg]
        parts = [(s["median_stt_ms"], C_TEAL, "speech-to-text"), (s["median_agent_ms"], C_VIOLET, "agent (LLM calls)"),
                 (s["median_tts_reply_ms"], C_MARIGOLD, "reply speech")]
        left = 0.0
        for ms, col, lab in parts:
            ax.barh(y, ms / 1000, 0.55, left=left, color=col, edgecolor=PAPER, linewidth=2, label=lab if y == ys[0] else None)
            left += ms / 1000
        ax.text(left + 0.15, y, f"{secs(s['median_ttfa_ms'])}", va="center", fontsize=11, color=INK)
    tgt = R["voice_roundtrip"]["target_ttfa_s"]
    ax.axvline(tgt, color=C_CORAL, lw=1.5)
    ax.text(tgt + 0.12, -0.62, f"target {tgt:.1f} s (not met)", color=C_CORAL, fontsize=10.5, va="center")
    ax.set_ylim(-0.85, len(order) - 0.5)
    ax.set_yticks(ys, [names[lg] for lg in order], fontsize=12)
    ax.set_xlabel("Seconds to first audio (median of each stage)")
    ax.set_title("Voice round trip, live")
    ax.set_xlim(0, 14.5)
    ax.legend(loc="upper center", bbox_to_anchor=(0.45, -0.17), ncol=3, fontsize=10)
    ax.grid(axis="y", visible=False)
    _src(fig, f"Sources: backend/reports/agent_suite.json ({gen(ag)}; held_out_before_fix), voice_roundtrip.json "
              f"({gen(R['voice_roundtrip'])}); {R['voice_roundtrip']['headline']['n_trips']} trips, synthetic speech in.", y=-0.13)
    return save(fig, "result_voice.png")


# ----------------------------------------------------------------------------- diagrams
def _canvas(w: float = 13.0, h: float = 6.2) -> tuple[plt.Figure, plt.Axes]:
    fig = plt.figure(figsize=(w, h))
    ax = fig.add_axes((0, 0, 1, 1))
    ax.set_xlim(0, 100 * w / h * 0.6)
    ax.set_ylim(0, 60)
    ax.axis("off")
    fig.patch.set_facecolor(PAPER)
    return fig, ax


def box(ax: plt.Axes, x: float, y: float, w: float, h: float, title: str, lines: list[str] | None = None,
        fc: str = CARD, ec: str = RULE, tc: str = INK, lc: str = MUTED, accent: str | None = None,
        ts: float = 13, ls: float = 10.5, align: str = "left") -> None:
    ax.add_patch(FancyBboxPatch((x, y), w, h, boxstyle="round,pad=0,rounding_size=1.2", fc=fc, ec=ec, lw=1.2))
    if accent:
        ax.add_patch(FancyBboxPatch((x, y + h - 0.9), w, 0.9, boxstyle="round,pad=0,rounding_size=0.45", fc=accent, ec="none"))
    tx = x + 1.2 if align == "left" else x + w / 2
    ha = "left" if align == "left" else "center"
    ax.text(tx, y + h - 2.3, title, fontsize=ts, fontweight="bold", color=tc, va="top", ha=ha)
    if lines:
        ax.text(tx, y + h - 2.3 - ts * 0.24, "\n".join(lines), fontsize=ls, color=lc, va="top", ha=ha, linespacing=1.45)


def arrow(ax: plt.Axes, x1: float, y1: float, x2: float, y2: float, color: str = MUTED, label: str | None = None,
          style: str = "-|>", lw: float = 1.6, ls: str = "-", lab_dx: float = 0.6, lab_dy: float = 0.0) -> None:
    ax.annotate("", xy=(x2, y2), xytext=(x1, y1),
                arrowprops={"arrowstyle": style, "color": color, "lw": lw, "linestyle": ls, "shrinkA": 0, "shrinkB": 0,
                            "mutation_scale": 16})
    if label:
        ax.text((x1 + x2) / 2 + lab_dx, (y1 + y2) / 2 + lab_dy, label, fontsize=9.5, color=color, va="center")


def diagram_system() -> Path:
    fig, ax = _canvas(13.0, 7.2)
    X = ax.get_xlim()[1]
    box(ax, 2, 52.5, X - 4, 6.5, "Browser  ·  React 18 + TypeScript + Vite  ·  Tailwind, Framer Motion, d3, i18next (bn / hi / kn / en)",
        ["Twin stage (replay clock, Two possible futures, Watch the twin learn) · Meal · What-if · Talk · Trust · Doctor (review queue, brief, lab upload)"],
        accent=MARIGOLD, ts=12, ls=10)
    arrow(ax, X - 5, 52.5, X - 5, 49.6)
    ax.text(X - 6, 51.05, "HTTPS (hosted demo: Caddy reverse proxy with TLS) · /api with JWT · signed audio URLs", fontsize=9.5,
            color=MUTED, va="center", ha="right")
    box(ax, 2, 40.3, X - 4, 9, "FastAPI  ·  pratifalan/api",
        ["auth, per-user ownership · replay clock (advance, reveal) · assimilate (unit, time, safety) · forecast · what-if · receipt",
         "meal + lab parsing (manifest ids, bounded uploads) · agent chat + confirm · doctor queue + review · FHIR · reports"],
        accent=INDIGO_2, ts=12, ls=9.6)
    cols = [
        ("Twin engine  ·  twin/", TEAL, ["model.py  7-state glucose ODE, RK4", "prior.py  EHR-conditioned prior",
                                         "calibrate.py  CEM multiple shooting", "filter.py  Liu-West particle filter",
                                         "hybrid.py  residual, conformal, events", "nbp.py  next-best-prick (EnSRF)",
                                         "engine.py  the Twin API + provenance"]),
        ("Perception  ·  perception/", MARIGOLD, ["meal.py  photo -> dish names + portions", "food.py  Indian food table -> carbs",
                                                  "lab.py  report -> fields, unit conversion", "     in code, collection dates",
                                                  "     -> covariate revision -> FHIR R4", "speech.py  STT / TTS adapters"]),
        ("Agent  ·  agent/", VIOLET, ["safety.py  one shared policy (manual,", "     chat, reveal); dose block, emergency",
                                      "orchestrator.py  router -> planner", "slots.py  numbers only via tool fields",
                                      "verifier.py  numeric + semantic check", "confirm before logging; audit trail"]),
        ("Providers  ·  providers/", CORAL, ["llm.py  one client, per-role model chains", "OpenRouter · OpenAI · DeepSeek",
                                             "Kimi · xAI · NanoGPT (fallback order)", "speech: OpenAI, Fish Audio,",
                                             "Sarvam adapter (no key yet)", "demo mode: recorded fixtures"]),
    ]
    cw = (X - 4 - 3 * 1.5) / 4
    for i, (t, acc, lines) in enumerate(cols):
        x = 2 + i * (cw + 1.5)
        box(ax, x, 13.5, cw, 24.5, t, lines, accent=acc, ts=12, ls=10)
        arrow(ax, x + cw / 2, 40.3, x + cw / 2, 38.2)
    stores = [("PostgreSQL 16", ["users · twin_events · replay clocks · audit_log", "fhir_resources · lab revisions · reviews (per user)"]),
              ("Parquet + artifacts/", ["persona series · prior fits · hybrid model", "persona states · manifest.json (hashes, membership)"]),
              ("JSON + FHIR R4", ["reports/*.json (Trust panel) · FHIR bundles", "fixtures/ (offline demo) · evidence receipts"])]
    sw = (X - 4 - 2 * 1.5) / 3
    for i, (t, lines) in enumerate(stores):
        x = 2 + i * (sw + 1.5)
        box(ax, x, 1, sw, 9.5, t, lines, fc="#F1ECE2", ts=12, ls=9.8)
    arrow(ax, X / 2, 13.5, X / 2, 10.7)
    return save(fig, "arch_system.png")


def diagram_forecast_flow() -> Path:
    fig, ax = _canvas(13.0, 6.4)
    steps = [
        ("1  Request", ["POST /twin/{pid}/forecast or what_if", "ladder level + confirmed dinner", "(UI or an agent tool call)"], MARIGOLD),
        ("2  Replay clock", ["per user + persona: day, time;", "Advance runs recorded meals and", "the ladder's pricks at their times"], TEAL),
        ("3  User events", ["readings: unit, observed_at, safety", "check, then Gaussian likelihood", "with adaptive prick inflation"], TEAL),
        ("4  Propagate", ["particles + scenario variants in one", "batched RK4 loop, common random", "numbers (paired futures)"], TEAL),
        ("5  Learned layer", ["LightGBM residual per horizon", "(monotone in carbs) on the", "mechanistic median"], VIOLET),
        ("6  Honest band", ["normalised split conformal by ladder", "x hours-since-reading x horizon;", "beyond 120 min: exploratory"], VIOLET),
        ("7  Event risk", ["P(>180 in 2 h): stacked + isotonic,", "with held-out reliability of its bin;", "P(<70): NOT validated"], VIOLET),
        ("8  Abstain check", ["far from training data or particle", "collapse -> 'not sure, take a", "reading'"], CORAL),
        ("9  Receipt + reply", ["provenance: model version, inputs", "revision, observations, replay time;", "drivers: physiology vs learned"], MARIGOLD),
    ]
    X = ax.get_xlim()[1]
    w, h = (X - 4 - 2 * 3) / 3, 15.5
    pos = []
    for i, (t, lines, acc) in enumerate(steps):
        r, c = divmod(i, 3)
        if r == 1:
            c = 2 - c  # snake layout
        x = 2 + c * (w + 3)
        y = 41 - r * (h + 4)
        box(ax, x, y, w, h, t, lines, accent=acc, ts=13, ls=11)
        pos.append((x, y))
    for i in range(len(pos) - 1):
        (x1, y1), (x2, y2) = pos[i], pos[i + 1]
        if abs(y1 - y2) < 1:
            if x2 > x1:
                arrow(ax, x1 + w, y1 + h / 2, x2, y2 + h / 2)
            else:
                arrow(ax, x1, y1 + h / 2, x2 + w, y2 + h / 2)
        else:
            arrow(ax, x1 + w / 2, y1, x2 + w / 2, y2 + h)
    return save(fig, "arch_forecast_flow.png")


def diagram_agent() -> Path:
    fig, ax = _canvas(13.0, 6.3)
    X = ax.get_xlim()[1]
    row = [
        ("Speech or text in", ["STT (OpenAI; Sarvam ready)", "or typed text; user sees", "the transcript"], MARIGOLD),
        ("1  Safety rules", ["deterministic, first: dose /", "medicine -> block; reading or", "symptoms -> shared policy"], CORAL),
        ("2  Router", ["rules first; small LLM only", "when unsure (intent + its", "own safety label)"], VIOLET),
        ("3  Planner tools", ["may call ONLY the Twin API:", "get_state · forecast · what_if", "next_best_prick · lookup_food"], VIOLET),
    ]
    row2 = [
        ("4  Slots", ["planner writes slots, not", "numbers; code fills them", "from named tool fields"], TEAL),
        ("5  Verifier", ["code: every number traced to", "this turn's tools; high/low", "risk bound to the right field"], TEAL),
        ("6  Output safety", ["re-check the final reply", "for dose or medicine advice"], CORAL),
        ("7  Localise + audit", ["bn / hi / kn / en, Indic digits", "re-verified; tool calls stored", "for the Why? drawer"], MARIGOLD),
    ]
    w, h = (X - 4 - 3 * 2.5) / 4, 15
    for i, (t, lines, acc) in enumerate(row):
        x = 2 + i * (w + 2.5)
        box(ax, x, 38, w, h, t, lines, accent=acc, ts=13, ls=10.5)
        if i:
            arrow(ax, x - 2.5, 45.5, x, 45.5)
    for i, (t, lines, acc) in enumerate(row2[::-1]):
        x = 2 + i * (w + 2.5)
        box(ax, x, 15, w, h, t, lines, accent=acc, ts=13, ls=10.5)
        if i:
            arrow(ax, x, 22.5, x - 2.5, 22.5)
    xr = 2 + 3 * (w + 2.5)
    arrow(ax, xr + w / 2, 38, xr + w / 2, 30.2)
    box(ax, 2 + 2 * (w + 2.5), 1.5, 2 * w + 2.5, 10, "Template fallback (also the whole offline demo)",
        ["reply built from tool outputs in 4 languages; always passes the verifier"], fc="#F1ECE2", ts=12, ls=10.5)
    arrow(ax, xr + w / 2, 15, xr + w / 2, 11.7, color=C_CORAL, ls="--", label="fails twice", lab_dx=0.8)
    box(ax, 2, 1.5, 2 * w + 2.5, 10, "Confirm before mutating the twin",
        ["log_reading / log_meal -> pending action -> 'Yes, add it' -> POST /api/agent/confirm"], fc="#F1ECE2", ts=12, ls=10.5)
    arrow(ax, 2 + w / 2, 15, 2 + w / 2, 11.7, color=C_TEAL, ls="--")
    xs1 = 2 + (w + 2.5) + w / 2
    arrow(ax, xs1, 38, xs1, 30.2, color=C_CORAL, ls="--", label="blocked / emergency:\nfixed message, spoken", lab_dx=0.8)
    ax.text(2, 57.3, "No LLM ever produces a glucose number: numbers come from the Twin API, words from the LLM.",
            fontsize=12.5, color=INK, fontweight="bold")
    return save(fig, "arch_agent_contract.png")


def diagram_eval(V: Values) -> Path:
    fig, ax = _canvas(13.0, 6.4)
    X = ax.get_xlim()[1]
    ax.text(2, 57.3, f"Nested person-wise {V['man_folds']}-fold cross-validation: no learned part ever sees the people it is scored on",
            fontsize=12.5, color=INK, fontweight="bold")
    w = (X - 4 - 2 * 2.5) / 3
    box(ax, 2, 27, w, 26, f"Outer fold k (by person)", [
        f"CGMacros: {V['sl_n_patients']} adults, split by person", "test people: about 9-10 per fold",
        "each: 5-day CGM calibration wear,", "then only the ladder's pricks", "(simulated, 5% glucometer error)",
        "scored against the hidden Dexcom CGM", f"{V['n_windows']} forecast windows per level"], accent=MARIGOLD, ts=13, ls=10.5)
    box(ax, 2 + w + 2.5, 27, w, 26, "Training people of fold k only", [
        "population prior (EHR -> parameters)", "stage-1 replays of training rows with", "priors that also exclude fold k",
        "inner folds BY PERSON for the hybrid:", "residual, conformal, classifier,", "isotonic, abstain; baselines too",
        "test people use a prior fitted on", "the training people only"], accent=TEAL, ts=13, ls=10.5)
    box(ax, 2 + 2 * (w + 2.5), 27, w, 26, "artifacts/manifest.json", [
        "membership of every fitted dependency", f"leakage audit: {V['man_checks']} checks {V['man_passed']}",
        "code hash + data hash -> cache keys", f"model version {V['man_version']}", "95% CIs: bootstrap over PEOPLE",
        "(ShanghaiT2DM: repeat recordings", "resampled together)"], accent=VIOLET, ts=13, ls=10.5)
    arrow(ax, 2 + w + 2.5, 40, 2 + w, 40, label=None)
    arrow(ax, 2 + 2 * w + 2.5, 40, 2 + 2 * w + 5, 40)
    box(ax, 2, 2, (X - 4 - 2.5) / 2, 21, "Production model behind the demo", [
        f"prior + hybrid fitted on {V['prod_people']} CGMacros people", f"excludes the persona sources {V['prod_excluded']},",
        "so the three demo personas are people", "the deployed model never saw", "(names and EHR narratives are invented)"],
        fc="#F1ECE2", ts=12.5, ls=10.5)
    box(ax, 2 + (X - 4 - 2.5) / 2 + 2.5, 2, (X - 4 - 2.5) / 2, 21, "Protocol change, old -> new (2 pricks/day)", [
        f"RMSE 60 min {V['pc_rmse_old']} -> {V['pc_rmse_new']} mg/dL ({V['pc_rmse_d']})",
        f"AUROC P(>180) {V['pc_auroc_old']} -> {V['pc_auroc_new']}",
        f"coverage {V['pc_cov_old']} -> {V['pc_cov_new']}; width {V['pc_w_old']} -> {V['pc_w_new']}",
        f"transfer gap {V['pc_gap_old']} -> {V['pc_gap_new']} mg/dL", "headline numbers barely moved"],
        fc="#F1ECE2", ts=12.5, ls=10.5)
    return save(fig, "arch_eval_protocol.png")


def diagram_engine() -> Path:
    fig, ax = _canvas(8.3, 4.95)
    X = ax.get_xlim()[1]
    layers = [
        ("Mechanistic layer", "7-state glucose-insulin ODE: own insulin secretion, 2-compartment gut,\n"
                              "exercise uptake, dawn effect, drift. RK4 at 5 min over particles.", TEAL),
        ("Personalisation", "EHR-conditioned prior -> CEM fit on the 5-day CGM wear -> Liu-West\n"
                            "particle filter; finger-pricks: Gaussian likelihood, adaptive inflation.", MARIGOLD),
        ("Learned correction", "LightGBM residual (monotone in carbs) · normalised split conformal\n"
                               "bands · stacked event classifier + isotonic · abstain.", VIOLET),
        ("Decisions", "Paired what-if (Two possible futures) on the same particles · next\n"
                      "best prick (EnSRF, experimental) · drivers: physiology vs learned.", CORAL),
    ]
    h = 12.6
    for i, (t, body, acc) in enumerate(layers):
        y = 60 - (i + 1) * (h + 2.0) + 1.0
        ax.add_patch(FancyBboxPatch((1, y), X - 2, h, boxstyle="round,pad=0,rounding_size=1.0", fc=CARD, ec=RULE, lw=1.2))
        ax.add_patch(FancyBboxPatch((1, y), 1.2, h, boxstyle="round,pad=0,rounding_size=0.5", fc=acc, ec="none"))
        ax.text(3.6, y + h - 2.2, t, fontsize=14, fontweight="bold", color=INK, va="top")
        ax.text(3.6, y + h - 6.0, body, fontsize=11.5, color=MUTED, va="top", linespacing=1.4)
    return save(fig, "engine_layers.png")


def diagram_idea() -> Path:
    fig, ax = _canvas(13.0, 3.6)
    X = ax.get_xlim()[1]
    w = (X - 4 - 2 * 4) / 3
    box(ax, 2, 5, w, 50, "Calibrate", ["an initial 5-day CGM wear", "+ EHR: age, sex, BMI, HbA1c", "(a clinic-lent sensor)"],
        accent=TEAL, ts=22, ls=17)
    box(ax, 2 + w + 4, 5, w, 50, "Keep it alive", ["about 2 finger-pricks a day", "Indian meals -> food table carbs", "activity; lab uploads"],
        accent=MARIGOLD, ts=22, ls=17)
    box(ax, 2 + 2 * (w + 4), 5, w, 50, "See, ask, learn", ["P(>180 mg/dL in the next 2 h)", "two possible futures, 90% band",
                                                          "watch it learn from a reading", "ask in bn / hi / kn / en"],
        accent=VIOLET, ts=22, ls=17)
    arrow(ax, 2 + w + 0.5, 30, 2 + w + 3.5, 30, lw=2.5)
    arrow(ax, 2 + 2 * w + 4.5, 30, 2 + 2 * w + 7.5, 30, lw=2.5)
    return save(fig, "idea_flow.png")


# ----------------------------------------------------------------------------- screenshots
# Real-backend screenshots (demo mode) taken at 1366 px (desktop) and 390 px (phone).
SHOT_PLAN = {  # target name: (source file name, crop box (left, top, right, bottom) or None)
    "stage_en.png": ("en-IN_1366__.png", (0, 64, 1366, 1062)),
    "whatif_en.png": ("en-IN_1366__whatif.png", (40, 88, 1320, 690)),
    "trust_en.png": ("en-IN_1366__trust.png", (40, 88, 1326, 1010)),
    "doctor_en.png": ("en-IN_1366__doctor.png", (40, 88, 1326, 575)),
    "meal_en.png": ("en-IN_1366__meal.png", (100, 88, 1266, 720)),
    "hi_mobile.png": ("hi-IN_390__.png", (0, 0, 390, 1100)),
    "kn_talk.png": ("kn-IN_1366__talk.png", (60, 64, 1320, 830)),
}


def screenshots(src: Path | None) -> None:
    SCREENS.mkdir(parents=True, exist_ok=True)
    for name, (srcname, crop) in SHOT_PLAN.items():
        dst = SCREENS / name
        if src is not None and (src / srcname).exists():
            im = Image.open(src / srcname).convert("RGB")
            if crop:
                im = im.crop((crop[0], crop[1], min(crop[2], im.width), min(crop[3], im.height)))
            im.save(dst, optimize=True)
        elif not dst.exists():
            print(f"[make_docs] screenshot {name} missing (pass --shots DIR)", file=sys.stderr)


# ----------------------------------------------------------------------------- pptx helpers
def rgb(h: str) -> RGBColor:
    return RGBColor.from_string(h.lstrip("#"))


def font_for(text: str) -> str:
    for ch in text:
        o = ord(ch)
        if 0x0900 <= o <= 0x097F:
            return INDIC_FONTS["deva"]
        if 0x0980 <= o <= 0x09FF:
            return INDIC_FONTS["beng"]
        if 0x0C80 <= o <= 0x0CFF:
            return INDIC_FONTS["knda"]
    return FONT


class Deck:
    W, H = Inches(13.333), Inches(7.5)

    def __init__(self, footer_name: str) -> None:
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = self.W, self.H
        self.footer_name = footer_name
        self.n = 0

    def slide(self, dark: bool = False, notes: str = "") -> Any:
        s = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = rgb(INDIGO if dark else PAPER)
        self.n += 1
        s._dark = dark  # type: ignore[attr-defined]
        if notes:
            s.notes_slide.notes_text_frame.text = notes
        self.text(s, 0.6, 7.05, 9, 0.3, self.footer_name, 10, MUTED if not dark else "#9CA0C8")
        self.text(s, 11.9, 7.05, 0.85, 0.3, str(self.n), 10, MUTED if not dark else "#9CA0C8", align="right")
        return s

    def text(self, s: Any, x: float, y: float, w: float, h: float, paras: Any, size: float = 18,
             color: str = INK, bold: bool = False, align: str = "left", anchor: str = "top",
             spacing: float = 1.1, after: float = 6) -> Any:
        tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = Inches(0.02)
        tf.margin_top = tf.margin_bottom = Inches(0.02)
        tf.vertical_anchor = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}[anchor]
        if isinstance(paras, str):
            paras = [paras]
        for i, para in enumerate(paras):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}[align]
            p.line_spacing = spacing
            p.space_after = Pt(after)
            runs = para if isinstance(para, list) else [(para, {})]
            for rt, st in runs:
                r = p.add_run()
                r.text = rt
                f = r.font
                f.name = st.get("font", font_for(rt))
                f.size = Pt(st.get("size", size))
                f.bold = st.get("bold", bold)
                f.italic = st.get("italic", False)
                f.color.rgb = rgb(st.get("color", color))
                # Indic scripts are "complex scripts": LibreOffice and PowerPoint pick the a:cs typeface for them.
                latin = r._r.get_or_add_rPr().find(qn("a:latin"))
                if latin is not None and f.name != FONT:
                    cs = OxmlElement("a:cs")
                    cs.set("typeface", f.name)
                    latin.addnext(cs)
        return tb

    def kicker(self, s: Any, text: str) -> None:
        self.text(s, 0.6, 0.35, 12, 0.35, text.upper(), 12, MARIGOLD if s._dark else C_MARIGOLD, bold=True)

    def title(self, s: Any, text: str, sub: str | None = None, size: float = 34) -> None:
        """One-line title: shrink the font (down to 22 pt) until the estimated width fits 12.1 in."""
        dark = s._dark
        size = max(22.0, min(size, 12.1 * 72 / (0.58 * max(len(text), 1))))
        self.text(s, 0.6, 0.7, 12.1, 0.9, text, size, PAPER if dark else INK, bold=True, spacing=1.0)
        if sub:
            self.text(s, 0.6, 1.5, 12.1, 0.6, sub, 16, "#C9CBE6" if dark else MUTED, spacing=1.05)

    def source(self, s: Any, text: str) -> None:
        self.text(s, 0.6, 6.68, 12.1, 0.35, text, 10, "#9CA0C8" if s._dark else MUTED)

    def rect(self, s: Any, x: float, y: float, w: float, h: float, fill: str, line: str | None = None,
             radius: float = 0.08) -> Any:
        sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        sh.adjustments[0] = radius
        sh.fill.solid()
        sh.fill.fore_color.rgb = rgb(fill)
        if line:
            sh.line.color.rgb = rgb(line)
            sh.line.width = Pt(1)
        else:
            sh.line.fill.background()
        sh.shadow.inherit = False
        return sh

    def tile(self, s: Any, x: float, y: float, w: float, h: float, big: str, label: str, accent: str,
             big_size: float = 30, label_size: float = 12.5) -> None:
        dark = s._dark
        self.rect(s, x, y, w, h, INDIGO_2 if dark else CARD, None if dark else RULE)
        self.rect(s, x, y + 0.15, 0.09, h - 0.3, accent, radius=0.5)
        self.text(s, x + 0.28, y + 0.1, w - 0.4, 0.7, big, big_size, PAPER if dark else INK, bold=True)
        self.text(s, x + 0.28, y + 0.12 + big_size / 72 * 1.3, w - 0.4, h - 0.8, label, label_size,
                  "#C9CBE6" if dark else MUTED, spacing=1.02, after=2)

    def card(self, s: Any, x: float, y: float, w: float, h: float, head: str, body: list[str], accent: str,
             head_size: float = 19, body_size: float = 15) -> None:
        dark = s._dark
        self.rect(s, x, y, w, h, INDIGO_2 if dark else CARD, None if dark else RULE)
        self.rect(s, x + 0.25, y + 0.25, 0.55, 0.08, accent, radius=0.5)
        self.text(s, x + 0.25, y + 0.42, w - 0.5, 0.55, head, head_size, PAPER if dark else INK, bold=True)
        self.text(s, x + 0.25, y + 0.42 + head_size / 72 * 1.6, w - 0.5, h - 1.0,
                  [[("•  ", {"color": accent}), (b, {})] for b in body], body_size,
                  "#D7D9EE" if dark else INK, spacing=1.04, after=6)

    def pic(self, s: Any, path: Path, x: float, y: float, w: float, h: float, border: bool = False) -> tuple[float, float, float, float]:
        with Image.open(path) as im:
            iw, ih = im.size
        scale = min(w / iw, h / ih)
        pw, ph = iw * scale, ih * scale
        px, py = x + (w - pw) / 2, y + (h - ph) / 2
        if border:
            self.rect(s, px - 0.04, py - 0.04, pw + 0.08, ph + 0.08, RULE, radius=0.02)
        s.shapes.add_picture(str(path), Inches(px), Inches(py), Emu(int(Inches(pw))), Emu(int(Inches(ph))))
        return px, py, pw, ph

    def save(self, path: Path) -> None:
        self.prs.save(str(path))


# ----------------------------------------------------------------------------- architecture deck
def build_architecture(V: Values, figs: dict[str, Path]) -> Path:
    d = Deck(f"{TEAM} · architecture · {INSTITUTION}")
    s = d.slide(notes="System overview. Every glucose number is produced by the twin engine; the LLM only perceives "
                      "(photos, speech, lab reports) and explains. Hosted demo: Caddy reverse proxy with TLS in front of the "
                      "docker compose stack (web, api, PostgreSQL).")
    d.kicker(s, "Architecture · 1 of 4")
    d.title(s, "One twin engine behind every screen and every sentence", size=30)
    d.pic(s, figs["arch_system"], 0.5, 1.45, 12.33, 5.25)
    d.source(s, f"Hosted: {HOSTED} (Caddy, TLS) · local: docker compose up with an empty .env -> http://localhost:8080 "
                "(offline demo mode, recorded fixtures).")

    s = d.slide(notes="Data flow of one forecast request, from the UI or from an agent tool call.")
    d.kicker(s, "Architecture · 2 of 4")
    d.title(s, "Data flow of a forecast", "Deterministic and cached: the same timestamped inputs give the same numbers and the same receipt",
            size=30)
    d.pic(s, figs["arch_forecast_flow"], 0.5, 2.1, 12.33, 4.5)
    d.source(s, f"Validated horizon: {V['validated_h']} min; 180-240 min drawn as exploratory "
                f"(backend/reports/sensor_ladder.json, exploratory_horizons, {V['gen_sensor_ladder']}).")

    s = d.slide(notes="Agent contract. The localiser is merged into the planner; the verifier re-runs on the localised "
                      "text (docs/DEVIATIONS.md #18). Numbers enter replies only through slots.")
    d.kicker(s, "Architecture · 3 of 4")
    d.title(s, "The agent contract", "rules → router → planner tools → slots → verifier → output safety → localiser → audit", size=30)
    d.pic(s, figs["arch_agent_contract"], 0.5, 2.0, 12.33, 4.65)
    d.source(s, f"Agent suite, {V['ag_n']} prompts in 4 languages: grounded-or-template {V['t_grounded_or_fallback_rate']} / "
                f"{V['l_grounded_or_fallback_rate']} (offline / live); unsafe advice {V['t_unsafe_advice_rate']} / "
                f"{V['l_unsafe_advice_rate']} (backend/reports/agent_suite.json, {V['gen_agent_suite']}).")

    s = d.slide(notes="Evaluation protocol after the external review: nested outer folds, inner folds and CIs by person, "
                      "membership manifest with automated leakage checks, cache keys from code and data hashes.")
    d.kicker(s, "Architecture · 4 of 4")
    d.title(s, "Evaluation protocol: nested folds and a membership manifest", size=30)
    d.pic(s, figs["arch_eval_protocol"], 0.5, 1.45, 12.33, 5.15)
    d.source(s, f"backend/artifacts/manifest.json (leakage_audit, membership); backend/reports/summary.json (protocol_change, "
                f"{V['gen_summary']}).")
    out = DOCS / "architecture.pptx"
    d.save(out)
    return out


# ----------------------------------------------------------------------------- presentation deck
def build_presentation(V: Values, figs: dict[str, Path]) -> Path:
    d = Deck(f"{TEAM} · {CHALLENGE}")
    light = "#E6E7F5"

    # 1 problem
    s = d.slide(dark=True, notes="Open with the person, not the model.")
    d.text(s, 0.6, 0.5, 9, 0.5, [[("PRATIFALAN  ·  ", {}), ("প্রতিফলন", {}), ("  ·  ", {}), ("प्रतिफलन", {}),
                                  ("  ·  ", {}), ("ಪ್ರತಿಫಲನ", {})]], 14, MARIGOLD, bold=True)
    d.text(s, 0.6, 1.0, 11.8, 1.6, "Reflect your health issues before they really occur", 40, PAPER, bold=True, spacing=0.95)
    d.text(s, 0.6, 2.85, 11.6, 0.5, "The problem", 17, MARIGOLD, bold=True)
    d.text(s, 0.6, 3.3, 11.9, 2.4, [
        "Most adults with Type 2 diabetes in India live between clinic visits with a glucometer and a few "
        "finger-pricks, not a continuous glucose monitor.",
        "The spike after a rice-heavy dinner happens when nobody is measuring, and the tools that could warn about it "
        "assume a daily sensor, Western food and English.",
    ], 21, light, spacing=1.08, after=12)
    d.text(s, 0.6, 5.95, 11.9, 0.6, f"Team {TEAM} · {MEMBER} · {INSTITUTION}", 15, PAPER, bold=True)
    d.text(s, 0.6, 6.4, 11.9, 0.4, f"Hosted demo: {HOSTED}", 13, "#9CA0C8")

    # 2 why CGM twins miss India
    s = d.slide(notes="Three gaps. Keep it to the gaps that the product answers.")
    d.kicker(s, "02 · Why CGM twins miss India")
    d.title(s, "Glucose twins are built for people who wear a sensor every day")
    cw = 3.85
    d.card(s, 0.6, 2.0, cw, 4.5, "Sensor", ["Most twins re-read a CGM every few minutes.",
                                            "A continuous sensor is a recurring cost most patients here do not carry.",
                                            "Finger-pricks are what people already have."], TEAL, body_size=17)
    d.card(s, 0.6 + cw + 0.3, 2.0, cw, 4.5, "Food", ["Thali, katori, roti count: not grams on a label.",
                                                     "No open Indian CGM dataset exists to learn from.",
                                                     f"A vision LLM's carb guesses barely tracked logs (Spearman {V['mp_rho']}, "
                                                     f"{V['mp_n']} photos)."], MARIGOLD, body_size=17)
    d.card(s, 0.6 + 2 * (cw + 0.3), 2.0, cw, 4.5, "Language", ["Advice apps speak English.",
                                                               "People ask in Bengali, Hindi, Kannada.",
                                                               "They need a short spoken answer with numbers they can trust."], VIOLET, body_size=17)
    d.source(s, f"backend/reports/meal_photo.json ({V['gen_meal_photo']}): US CGMacros meal photos; see slide 10 for the forecast cost.")

    # 3 idea
    s = d.slide(notes="Showcase line from the review: a sensor-light glucose twin with Indian meal inputs, multilingual "
                      "interaction, and visible uncertainty. Always say the 5-day CGM calibration out loud.")
    d.kicker(s, "03 · The idea")
    d.title(s, "A sensor-light glucose twin with Indian meal inputs, multilingual interaction and visible uncertainty",
            size=26)
    d.text(s, 0.6, 1.6, 12.1, 0.6, [[("Target outcome: ", {"bold": True}),
           ("P(glucose > 180 mg/dL within 2 h) after a meal, for adults with T2D, with a 90% band. "
            "P(<70) is shown but labelled not validated.", {})]], 16, MUTED)
    d.pic(s, figs["idea_flow"], 0.5, 2.35, 12.33, 4.2)
    d.source(s, f"Tested setup: a 5-day CGM calibration wear, then finger-pricks. P(<70) not validated: {V['low_windows']} low "
                f"windows from {V['low_patients']} of {V['sl_n_patients']} people (backend/reports/sensor_ladder.json, low_events).")

    # 4 demo screenshots
    s = d.slide(notes="Screenshots of the real app against the real API (demo mode, recorded fixtures). "
                      f"Hosted at {HOSTED}; login TestUser / TestUser11 (shown on the sign-in page).")
    d.kicker(s, "04 · Live demo")
    d.title(s, "Tonight, with your twin", f"{HOSTED} · personas are synthetic composites the production model never saw", size=32)
    top = 2.05
    px, py, pw, ph = d.pic(s, SCREENS / "stage_en.png", 0.6, top, 6.3, 4.25, border=True)
    d.text(s, px, py + ph + 0.06, pw, 0.3, "Twin stage: replay clock, risk with held-out reliability, Watch the twin learn",
           12, MUTED, align="center")
    px, py, pw, ph = d.pic(s, SCREENS / "hi_mobile.png", 7.1, top, 1.55, 4.25, border=True)
    d.text(s, px - 0.25, py + ph + 0.06, pw + 0.5, 0.3, "Hindi, 390 px", 12, MUTED, align="center")
    px, py, pw, ph = d.pic(s, SCREENS / "doctor_en.png", 8.9, top, 3.85, 1.6, border=True)
    d.text(s, px, py + ph + 0.04, pw, 0.3, "Clinician review queue", 12, MUTED, align="center")
    px, py, pw, ph = d.pic(s, SCREENS / "kn_talk.png", 8.9, top + 2.05, 3.85, 2.2, border=True)
    d.text(s, px, py + ph + 0.06, pw, 0.3, "Talk in Kannada", 12, MUTED, align="center")

    # 5 twin engine
    s = d.slide(notes="Four layers. Honest point: at 2 pricks a day the hybrid is level with a LightGBM model on sparse "
                      "features (paired CI includes zero). It clearly beats the mechanistic layer alone, the personal average "
                      "and persistence; its added value is calibrated bands, paired futures and explanations.")
    d.kicker(s, "05 · Twin engine")
    d.title(s, "Physiology first, learning second, honesty last", size=32)
    d.pic(s, figs["engine_layers"], 0.5, 1.45, 8.1, 4.95)
    d.text(s, 8.85, 1.5, 3.95, 0.6, "60-min RMSE, 2 pricks/day; paired difference vs the hybrid (95% CI)", 13, INK, bold=True,
           spacing=1.0)
    rows = [("Hybrid twin", V["b_full_hybrid_p2"], "", C_MARIGOLD),
            ("LightGBM, sparse", V["b_lightgbm_sparse_only_p2"], f"{V['bp_lightgbm_sparse_only_p2']} ({V['bp_lightgbm_sparse_only_lo_p2']} to {V['bp_lightgbm_sparse_only_hi_p2']})", INK),
            ("Mechanistic only", V["b_mechanistic_only_p2"], f"{V['bp_mechanistic_only_p2']} ({V['bp_mechanistic_only_lo_p2']} to {V['bp_mechanistic_only_hi_p2']})", INK),
            ("Personal average", V["b_personal_average_p2"], f"{V['bp_personal_average_p2']} ({V['bp_personal_average_lo_p2']} to {V['bp_personal_average_hi_p2']})", INK),
            ("Persistence", V["b_persistence_p2"], f"{V['bp_persistence_p2']} ({V['bp_persistence_lo_p2']} to {V['bp_persistence_hi_p2']})", INK)]
    for i, (lab, v, dd, col) in enumerate(rows):
        y = 2.25 + i * 0.66
        d.text(s, 8.85, y, 2.5, 0.35, lab, 14, col, bold=(i == 0))
        d.text(s, 11.75, y, 1.0, 0.35, v, 16, col, bold=True, align="right")
        if dd:
            d.text(s, 8.85, y + 0.3, 3.9, 0.3, dd, 11, MUTED)
    d.text(s, 8.85, 5.6, 3.95, 0.9, f"Level with LightGBM at 2 pricks/day (CI includes 0). Raw particle P(>180) ECE "
                                    f"{V['ece_uncal_p2']} → {V['ece_cal_p2']} after stacking + isotonic.", 12, MUTED, spacing=1.02)
    d.source(s, f"mg/dL; backend/reports/baselines.json (paired_vs_full_hybrid_60min, {V['gen_baselines']}), calibration_curve.json. "
                "Same forecast windows for every method; nested person-wise CV.")

    # 6 sensor ladder
    s = d.slide(notes="Lead with the measured trade-off. The by-hours-since-reading table is confounded by clock time "
                      "(pricks at 07:00 and 22:00), so the controlled evidence is the paired band width at identical origins.")
    d.kicker(s, "06 · Sensor-ladder result (Experiment 1)")
    d.title(s, f"2 pricks a day after 5 CGM days: 60-min error {V['rmse60_p2']} mg/dL, band covers {V['cov60_p2']}",
            f"Full CGM {V['rmse60_full']} · no glucose data {V['rmse60_p0']} · band width {V['width60_p2']} mg/dL · "
            f"AUROC P(>180) {V['auroc_p2']} ({V['auroc_lo_p2']}–{V['auroc_hi_p2']})", size=28)
    d.pic(s, figs["result_sensor_ladder"], 0.45, 2.1, 12.43, 3.95)
    d.text(s, 0.6, 6.08, 12.1, 0.6, [[("Living uncertainty, controlled: ", {"bold": True}),
           (f"within 90 min of a prick the band is {V['pbw_near_d60']} mg/dL vs no data at 60 min ({V['pbw_near_lo60']} to "
            f"{V['pbw_near_hi60']}); 8 h or more later it is {V['pbw_far_d60']} ({V['pbw_far_lo60']} to {V['pbw_far_hi60']}). "
            f"Recent prick: 60-min error {V['rp_diff']} mg/dL ({V['rp_lo']} to {V['rp_hi']}).", {})]], 13.5, INK)

    # 7 NBP
    s = d.slide(notes="Report the null as it is. The feature stays in the app as an experimental information-gathering "
                      "suggestion, not an accuracy promise.")
    d.kicker(s, "07 · Next-best-prick result (Experiment 4)")
    d.title(s, "Choosing when to prick did not beat fixed clinic times", "An honest null result, kept in the app only as an experimental suggestion")
    d.pic(s, figs["result_nbp"], 0.5, 2.2, 8.4, 3.95)
    d.tile(s, 9.25, 2.3, 3.5, 1.15, V["nbp_twin_rmse60"], "60-min RMSE, twin-recommended times", C_VIOLET)
    d.tile(s, 9.25, 3.6, 3.5, 1.15, V["nbp_fixed_rmse60"], "60-min RMSE, fixed 07:00 & 22:00", C_GRAY)
    d.tile(s, 9.25, 4.9, 3.5, 1.15, V["nbp_random_rmse60"], "60-min RMSE, random times", C_GRAY)
    d.source(s, f"backend/reports/nbp_value.json ({V['gen_nbp_value']}); planner: ensemble square-root Kalman variance reduction over 12 h; "
                "2 pricks/day in every arm.")

    # 8 trust and fairness
    s = d.slide(notes="Coverage with width, calibration, subgroups, transfer with people versus recordings, and the protocol "
                      "change. Headline numbers barely moved after removing the indirect leakage path; that is good news.")
    d.kicker(s, "08 · Trust and fairness (Experiments 5, 6, 8)")
    d.title(s, f"Calibrated and covered; transfer to Shanghai costs {V['gap_sh']} mg/dL", size=30)
    d.pic(s, figs["result_trust"], 0.45, 1.5, 12.43, 3.75)
    d.tile(s, 0.6, 5.35, 3.9, 1.25, f"{V['cov60_p2']} · {V['width60_p2']}", "coverage · width (mg/dL) of the 90% band at 60 min, 2 pricks/day",
           C_TEAL, big_size=24, label_size=11.5)
    d.tile(s, 4.7, 5.35, 3.9, 1.25, f"{V['pc_rmse_old']} → {V['pc_rmse_new']}",
           f"60-min RMSE, old → new protocol; AUROC {V['pc_auroc_old']} → {V['pc_auroc_new']}; gap {V['pc_gap_old']} → {V['pc_gap_new']}",
           C_MARIGOLD, big_size=24, label_size=11.5)
    d.tile(s, 8.8, 5.35, 3.95, 1.25, f"{V['man_checks']} checks",
           f"leakage audit {V['man_passed']}; insulin users RMSE {V['ins_rmse60']} ({V['ins_people']} people, {V['ins_rec']} recordings)",
           C_VIOLET, big_size=24, label_size=11.5)

    # 9 voice
    s = d.slide(notes="Safety targets met on the current suite, including the 24 new held-out v3 prompts; their pre-fix "
                      "misses are disclosed. Voice latency target is not met; streaming TTS is roadmap.")
    d.kicker(s, "09 · Voice in four languages")
    d.title(s, f"Grounded and safe on {V['ag_n']} prompts; first audio after {V['v_ttfa']}", size=30)
    d.pic(s, figs["result_voice"], 0.45, 1.45, 12.43, 4.55)
    d.text(s, 0.6, 6.05, 12.1, 0.6, [[("Disclosed: ", {"bold": True, "color": C_CORAL}),
           (f"held-out prompts before the fix: template emergency recall {V['hb_t_emerg']}, dose block {V['hb_t_block']} "
            f"({V['hb_n_fail']} misses, then fixed). {V['v_target']} first-audio target met on {V['v_share']} of {V['v_trips']} trips; "
            "streaming speech is roadmap. Readings and meals said by voice need a confirm tap.", {})]], 13, INK)

    # 10 safety and limits
    s = d.slide(notes="Two columns: what the design guarantees, and what the evidence does not show yet.")
    d.kicker(s, "10 · Safety and limits")
    d.title(s, "Decision support, not a medical device")
    d.card(s, 0.6, 1.65, 5.95, 4.95, "Built in", [
        "One safety policy for manual entry, chat and replay reveal: a reading of 45 → very-low emergency guidance (108).",
        "No dose or medicine advice; numbers only via slots from twin tools.",
        "Confirm before voice/chat logs a reading or meal; per-user ownership checks; manifest-only sample ids.",
        f"P(<70) not validated; beyond {V['validated_h']} min drawn as exploratory.",
        f"Abstains when unsure ({V['abstain_p2']} of forecasts at 2 pricks/day).",
    ], TEAL, body_size=14)
    d.card(s, 6.78, 1.65, 5.95, 4.95, "Not proven yet", [
        "No Indian CGM data; trained on US and Chinese adults.",
        "Needs a 5-day CGM wear first; pricks simulated except the ShanghaiT2DM arm.",
        f"Photo carbs worsen 60-min error by {V['mpf_diff60']} mg/dL; so dishes are editable chips + a {V['food_n']}-dish Indian table.",
        f"Insulin is not modelled (RMSE {V['ins_rmse60']} in insulin users).",
        "No clinical validation; no prospective study yet.",
    ], CORAL, body_size=14)
    d.source(s, "backend/reports/sensor_ladder.json, meal_photo.json, transfer.json, summary.json (limitations); "
                "backend/pratifalan/agent/safety.py (shared policy).")

    # 11 roadmap
    s = d.slide(notes="Roadmap is ordered by what would change the evidence, not by features.")
    d.kicker(s, "11 · Roadmap: India pilot")
    d.title(s, "From open-data evidence to evidence from Indian clinics")
    cw = 3.85
    d.card(s, 0.6, 1.8, cw, 4.75, "Next 3 months", ["Streaming speech to reach the 2.5 s first-audio target.",
                                                    "Real-speaker tests in four languages; native-speaker review of translations.",
                                                    "Independent clinician review of the safety suite."], MARIGOLD)
    d.card(s, 0.6 + cw + 0.3, 1.8, cw, 4.75, "Pilot", ["Prospective study with ethics approval, CTRI-registered.",
                                                       "Indian CGM collection: blinded CGM as reference, real pricks, thali photos.",
                                                       "Re-run every experiment on Indian data; model insulin explicitly."], TEAL)
    d.card(s, 0.6 + 2 * (cw + 0.3), 1.8, cw, 4.75, "Scale", ["ABDM integration: ABHA linkage, consent manager, FHIR R4 exchange.",
                                                             "CDSCO SaMD pathway if claims move beyond decision support.",
                                                             "Clinic-lent sensors for the 5-day calibration wear."], VIOLET)

    # 12 ask
    s = d.slide(dark=True, notes="End with the ask and the links.")
    d.kicker(s, "12 · The ask")
    d.title(s, "Help us earn the evidence", size=40)
    d.text(s, 0.6, 1.95, 11.9, 3.6, [
        [("1  ", {"color": MARIGOLD, "bold": True}), ("A clinical partner for a consented, ethics-approved pilot in India.", {})],
        [("2  ", {"color": MARIGOLD, "bold": True}), ("Access to Indian CGM and meal data under DPDP-compliant consent.", {})],
        [("3  ", {"color": MARIGOLD, "bold": True}), ("Mentoring on CDSCO software-as-a-medical-device and ABDM integration.", {})],
        [("4  ", {"color": MARIGOLD, "bold": True}), ("Speech and LLM credits to test with real speakers in four languages.", {})],
    ], 22, light, spacing=1.1, after=14)
    d.text(s, 0.6, 5.55, 11.9, 1.1, [f"Try it: {HOSTED}  (login TestUser / TestUser11)",
                                     f"Team {TEAM} · {MEMBER} · {INSTITUTION}",
                                     "Every number in this deck is generated from backend/reports/*.json by scripts/make_docs.py."],
           14, "#9CA0C8", after=4)
    out = DOCS / "presentation.pptx"
    d.save(out)
    if d.n != 12:
        sys.exit(f"[make_docs] presentation has {d.n} slides, expected 12")
    return out


# ----------------------------------------------------------------------------- README blocks
def readme_blocks(R: dict, V: Values) -> dict[str, str]:
    keyfacts = f"""| What | Result | Report |
|---|---|---|
| 60-min forecast error, 2 finger-pricks/day after a 5-day CGM calibration | RMSE {V['rmse60_p2']} mg/dL (95% CI {V['rmse60_lo_p2']}–{V['rmse60_hi_p2']}); full CGM {V['rmse60_full']}; no glucose data after calibration {V['rmse60_p0']} | `sensor_ladder.json` |
| 90% band, 2 pricks/day, 60 min | covers {V['cov60_p2']} ({V['cov60_lo_p2']}–{V['cov60_hi_p2']}) with width {V['width60_p2']} mg/dL; at meal starts only {V['meal_cov60_p2']} | `sensor_ladder.json` |
| P(>180 mg/dL within 2 h) | AUROC {V['auroc_p2']} ({V['auroc_lo_p2']}–{V['auroc_hi_p2']}), ECE {V['ece_p2']} (raw particles {V['ece_uncal_p2']}) | `sensor_ladder.json`, `calibration_curve.json` |
| P(<70 mg/dL within 2 h) | **Not validated**: {V['low_windows']} low windows from {V['low_patients']} of {V['sl_n_patients']} people | `sensor_ladder.json` (`low_events`) |
| Hybrid vs LightGBM on sparse features (2 pricks/day, paired) | {V['bp_lightgbm_sparse_only_p2']} mg/dL ({V['bp_lightgbm_sparse_only_lo_p2']} to {V['bp_lightgbm_sparse_only_hi_p2']}): **level** | `baselines.json` |
| Hybrid vs mechanistic only / personal average / persistence | {V['bp_mechanistic_only_p2']} / {V['bp_personal_average_p2']} / {V['bp_persistence_p2']} mg/dL, all CIs exclude 0 | `baselines.json` |
| Value of a recent finger-prick | 60-min RMSE {V['rp_diff']} mg/dL ({V['rp_lo']} to {V['rp_hi']}), {V['rp_n']} windows, {V['rp_people']} people | `sensor_ladder.json` (`recent_prick_value`) |
| Living uncertainty, paired at identical origins | band {V['pbw_near_d60']} mg/dL vs no data within 90 min of a prick; {V['pbw_far_d60']} mg/dL 8 h or more after one (60 min) | `sensor_ladder.json` (`paired_band_width`) |
| Next-best-prick timing (experimental) | **No significant gain**: twin vs fixed {V['nbp_tf_mean']} mg/dL ({V['nbp_tf_lo']} to {V['nbp_tf_hi']}), {V['nbp_tf_improved']}/{V['nbp_tf_n']} people improved | `nbp_value.json` |
| Transfer CGMacros → ShanghaiT2DM | 60-min RMSE {V['sh_tr_rmse60']} vs {V['sh_in_rmse60']} in-domain (gap {V['gap_sh']}); real finger-pricks {V['sh_tr_real_rmse60']} vs {V['sh_in_real_rmse60']}; {V['sh_people']} people, {V['sh_recordings']} recordings | `transfer.json` |
| Insulin users (ShanghaiT2DM, transfer arm) | 60-min RMSE {V['ins_rmse60']} ({V['ins_people']} people) vs {V['noins_rmse60']} without insulin ({V['noins_people']} people) | `transfer.json` |
| Evaluation protocol change (nested folds) | 2-prick 60-min RMSE {V['pc_rmse_old']} → {V['pc_rmse_new']}; transfer gap {V['pc_gap_old']} → {V['pc_gap_new']}; {V['man_checks']} leakage checks {V['man_passed']} | `summary.json`, `artifacts/manifest.json` |
| Meal-photo carbs (Experiment 7) | vision-LLM carbs vs logs: MAE {V['mp_mae']} g, Spearman {V['mp_rho']}; 60-min RMSE {V['mpf_diff60']} mg/dL ({V['mpf_lo60']}–{V['mpf_hi60']}) when used | `meal_photo.json` |
| Agent suite ({V['ag_n']} prompts, 4 languages; {V['ho_n']} held out) | unsafe advice {V['t_unsafe_advice_rate']} / {V['l_unsafe_advice_rate']}, emergency recall {V['t_emergency_recall']} / {V['l_emergency_recall']} (offline / live); held-out before fix: template emergency recall {V['hb_t_emerg']}, dose block {V['hb_t_block']} | `agent_suite.json` |
| Voice round trip (live) | median first audio after {V['v_ttfa']}; **{V['v_target']} target not met** ({V['v_share']} of {V['v_trips']} trips) | `voice_roundtrip.json` |
"""
    protocol = f"""| Headline (summary.json `protocol_change`) | Before (indirect leakage path) | After (nested) | Change |
|---|---|---|---|
| 60-min RMSE, 2 pricks/day (mg/dL) | {V['pc_rmse_old']} | {V['pc_rmse_new']} | {V['pc_rmse_d']} |
| AUROC P(>180), 2 pricks/day | {V['pc_auroc_old']} | {V['pc_auroc_new']} | {V['pc_auroc_d']} |
| 90% coverage at 60 min, 2 pricks/day | {V['pc_cov_old']} | {V['pc_cov_new']} | {V['pc_cov_d']} |
| 90% band width at 60 min (mg/dL) | {V['pc_w_old']} | {V['pc_w_new']} | {V['pc_w_d']} |
| 60-min RMSE, full CGM (mg/dL) | {V['pc_full_old']} | {V['pc_full_new']} | {V['pc_full_d']} |
| Transfer gap into ShanghaiT2DM, 2 pricks/day (mg/dL) | {V['pc_gap_old']} | {V['pc_gap_new']} | {V['pc_gap_d']} |

Leakage audit in `backend/artifacts/manifest.json`: {V['man_checks']} automated checks, {V['man_passed']}; model version
`{V['man_version']}`. Old numbers: `backend/reports/archive/pre_nested_protocol/headline.json`."""

    evidence = f"""**Validated on real, open data (retrospective, person-held-out):**

- 30–120-minute glucose forecasts, 90% bands and P(>180 within 2 h) on **CGMacros** ({V['sl_n_patients']} US adults: T2D,
  prediabetes and healthy) under nested person-wise cross-validation, scored against the hidden Dexcom G6 CGM:
  2-prick RMSE {V['rmse60_p2']} mg/dL at 60 min, coverage {V['cov60_p2']} at width {V['width60_p2']} mg/dL, AUROC {V['auroc_p2']}
  (`sensor_ladder.json`).
- Calibration of P(>180) (ECE {V['ece_cal_p2']}; held-out reliability per probability bin is shown next to the number in
  the app) and coverage by sex, age, BMI (Asian cut-offs) and HbA1c: {V['sg_n_flagged']} of {V['sg_n_groups']} subgroups flagged, but
  the smallest has {V['sg_small_n']} people (`calibration_curve.json`, `subgroups.json`).
- Transfer between CGMacros and **ShanghaiT2DM** ({V['sh_people']} people, {V['sh_recordings']} recordings), including the only arm
  with **real finger-prick** readings: gap {V['gap_sh_real']} mg/dL at 60 min (`transfer.json`).
- The controlled effect of a recent reading: lower 60-min error ({V['rp_diff']} mg/dL, CI {V['rp_lo']} to {V['rp_hi']}) and a
  narrower band at identical origins ({V['pbw_near_d60']} mg/dL, CI {V['pbw_near_lo60']} to {V['pbw_near_hi60']})
  (`sensor_ladder.json`).
- The cost of photo-estimated carbohydrate on forecasts, using real CGMacros meal photos (US meals): {V['mpf_diff60']} mg/dL at
  60 min (`meal_photo.json`).

**Synthetic or simulated:**

- The three demo personas are **synthetic composites**: real CGMacros trajectories ({V['prod_excluded']}) with invented Indian names
  and EHR narratives. The production model excludes these three people, so they are unseen.
- Finger-pricks in all CGMacros experiments are **simulated** from the hidden CGM with 5% glucometer error.
- The replay clock replays recorded data; "Reveal the next recorded reading" shows the dataset's hidden reference CGM.
- Sample lab reports are synthetic; sample thali photos are illustrative (credited). Voice tests use synthetic speech.
- The 90-day outlook is a GMI-based projection, not an HbA1c prediction.

**Not proven:**

- **Added accuracy over simple machine learning at 2 pricks/day**: the hybrid is level with LightGBM on sparse features
  ({V['bp_lightgbm_sparse_only_p2']} mg/dL, CI {V['bp_lightgbm_sparse_only_lo_p2']} to {V['bp_lightgbm_sparse_only_hi_p2']}). Its
  added value is calibrated bands, paired scenarios and explanations.
- **Sensor-light benefit on average**: 2 pricks/day ({V['rmse60_p2']}) is close to no glucose data after calibration
  ({V['rmse60_p0']}); the gain is concentrated right after a prick. The 5-day CGM wear is part of the tested setup.
- **Next-best-prick timing**: no significant gain over fixed or random times (null result).
- **P(<70 within 2 h)**: not validated. **Beyond 2 hours**: exploratory (180-min RMSE {V['x180_rmse_p2']}, coverage
  {V['x180_cov_p2']}).
- **Two possible futures**: model simulations, not proven causal effects of a portion change or a walk.
- **India**: no open Indian CGM dataset; Indian dish recognition checked on {V['mpi_n']} sample photos only (a smoke test).
- **Insulin users** (RMSE {V['ins_rmse60']}) and **voice latency** (median first audio {V['v_ttfa']} vs a {V['v_target']} target).
- **No clinical validation** of any kind and no prospective study."""

    rp = R["sensor_ladder"]["recent_prick_value"]
    pb = R["sensor_ladder"]["paired_band_width"]
    bins = R["sensor_ladder"]["by_hours_since_reading"]
    hs = ("30", "60", "90", "120")
    L = R["sensor_ladder"]["levels"]
    B2 = next(r for r in R["baselines"]["rows"] if r["level"] == "2")["paired_vs_full_hybrid_60min"]
    blabel = {"persistence": "Persistence", "population_time_of_day": "Population time-of-day curve",
              "personal_average": "Personal average", "lightgbm_sparse_only": "LightGBM on sparse features",
              "mechanistic_only": "Mechanistic twin only"}
    ag = R["agent_suite"]
    extra = [
        "### Hybrid vs baselines, paired at 60 min (Experiment 3)",
        "",
        "| 2 pricks/day: baseline | Baseline RMSE | Hybrid minus baseline (mg/dL) | 95% CI | CI excludes 0 |",
        "|---|---|---|---|---|",
        *[f"| {blabel.get(k, k)} | {f1(p['rmse60_baseline'])} | {sf2(p['difference'])} | {m2(p['difference_ci95'][0])} to "
          f"{m2(p['difference_ci95'][1])} | {'yes' if p['ci_excludes_zero'] else '**no**'} |" for k, p in B2.items()],
        "",
        f"Negative = hybrid better; same {B2['persistence']['n_forecast_windows']:,} windows, {B2['persistence']['n_people']} people; "
        "CI by bootstrap over people.",
        "",
        f"<sub>Source: `backend/reports/baselines.json`, generated {V['gen_baselines']}.</sub>",
        "",
        "### Value of a recent finger-prick and paired band width (Experiment 1, follow-up)",
        "",
        f"_{rp['definition']}_",
        "",
        "| Horizon | RMSE, prick in the last 90 min | RMSE, no glucose data | Band width, 2 pricks − none, ≤ 90 min after a prick (95% CI) | Band width, 2 pricks − none, ≥ 8 h after a prick (95% CI) |",
        "|---|---|---|---|---|",
        *[f"| {h} min | {f1(rp['rmse_with_prick'][h])} | {f1(rp['rmse_without'][h])} | "
          f"{sf1(pb['recent_prick']['by_horizon'][h]['paired_difference'])} ({m1(pb['recent_prick']['by_horizon'][h]['paired_difference_ci95'][0])} to "
          f"{m1(pb['recent_prick']['by_horizon'][h]['paired_difference_ci95'][1])}) | "
          f"{sf1(pb['control_far_from_prick']['by_horizon'][h]['paired_difference'])} ({m1(pb['control_far_from_prick']['by_horizon'][h]['paired_difference_ci95'][0])} to "
          f"{m1(pb['control_far_from_prick']['by_horizon'][h]['paired_difference_ci95'][1])}) |" for h in hs],
        "",
        f"60-min RMSE difference: {V['rp_diff']} mg/dL (95% CI {V['rp_lo']} to {V['rp_hi']}); {V['rp_n']} windows, {V['rp_people']} people. "
        f"Band widths are paired at identical forecast origins (same person, same time), so clock time and meals are held fixed: "
        f"after a reading the band is narrower than with no data; {V['pbw_far_subset']} it is wider (coverage {V['pbw_far_covw60']} "
        f"vs {V['pbw_far_covo60']} at 60 min).",
        "",
        "| Hours since last reading (descriptive, confounded by clock time) | Windows | RMSE 60 min | 90% coverage (60 min) | 90% band width (60 min) |",
        "|---|---|---|---|---|",
        *[f"| {hours_label(b['hours_since_reading'])} | {b['n_forecast_windows']:,} | {f1(b['rmse60'])} | {pct1(b['coverage90_60'])} | "
          f"{f1(b['width90_60'])} |" for b in bins],
        "",
        R["sensor_ladder"]["by_hours_since_reading_note"],
        "",
        f"<sub>Source: `backend/reports/sensor_ladder.json`, generated {V['gen_sensor_ladder']}.</sub>",
        "",
        "### Exploratory horizons beyond 2 hours",
        "",
        "| Glucose data | RMSE 180 min | 90% coverage 180 min | RMSE 240 min | 90% coverage 240 min |",
        "|---|---|---|---|---|",
        *[f"| {lv['label']} | {f1(lv['exploratory_horizons']['horizons']['180']['rmse'])} | "
          f"{pct1(lv['exploratory_horizons']['horizons']['180']['coverage90'])} | {f1(lv['exploratory_horizons']['horizons']['240']['rmse'])} | "
          f"{pct1(lv['exploratory_horizons']['horizons']['240']['coverage90'])} |" for lv in L],
        "",
        L[0]["exploratory_horizons"]["note"],
        "",
        "### Clarke error grid (Experiment 8)",
        "",
        "| Glucose data | Zones A+B, virtual CGM | Zones A+B, 60-min forecast | Zone A, 60-min forecast |",
        "|---|---|---|---|",
        *[f"| {V['label_' + t]} | {V['clarke_ab_' + t]} | {V['eg_f60_ab_' + t]} | {V['eg_f60_a_' + t]} |"
          for t in ("full", "p4", "p2", "p1", "p0")],
        "",
        f"<sub>Source: `backend/reports/error_grid.json`, generated {V['gen_error_grid']}.</sub>",
        "",
        "### Meal-photo carbohydrate error (Experiment 7)",
        "",
        f"{V['mp_n']} CGMacros (US) meal photos from {V['mp_people']} participants, vision model `{V['mp_model']}`, "
        "compared with the participant's logged carbohydrate. 95% CIs by participant bootstrap.",
        "",
        "| Metric | Value |",
        "|---|---|",
        f"| Mean absolute error | {V['mp_mae']} g ({V['mp_mae_lo']}–{V['mp_mae_hi']}) |",
        f"| Median absolute error | {V['mp_medae']} g |",
        f"| Median absolute % error | {V['mp_medape']} |",
        f"| Bias (photo − logged) | {V['mp_bias']} g ({V['mp_bias_lo']} to {V['mp_bias_hi']}); limits of agreement {V['mp_loa_lo']} to {V['mp_loa_hi']} g |",
        f"| Spearman correlation | {V['mp_rho']} ({V['mp_rho_lo']} to {V['mp_rho_hi']}) |",
        f"| Within ±20% | {V['mp_w20']} |",
        f"| Post hoc, without {V['mpx_excl']} repeated drink logs ({V['mpx_n']} photos) | Spearman {V['mpx_rho']} ({V['mpx_rho_lo']}–{V['mpx_rho_hi']}) |",
        f"| Forecast impact, 60-min RMSE ({V['mpf_n']} paired windows) | {V['mpf_logged60']} logged → {V['mpf_photo60']} photo carbs ({V['mpf_diff60']}, 95% CI {V['mpf_lo60']}–{V['mpf_hi60']}) |",
        f"| Forecast impact at meal starts ({V['mpf_meal_n']} windows) | {V['mpf_meal_diff60']} mg/dL ({V['mpf_meal_lo60']}–{V['mpf_meal_hi60']}); 90% coverage {V['mpf_mcov_logged']} → {V['mpf_mcov_photo']} |",
        f"| Indian sample photos (smoke test, {V['mpi_n']} photos, app pipeline) | dish detection recall {V['mpi_recall']}, precision {V['mpi_precision']} |",
        "",
        "Consequence in the app: the LLM only names dishes and household portions as editable chips; carbohydrate always comes "
        f"from the Indian food table ({V['food_n']} dishes from INDB 2024, USDA FoodData Central and IFCT 2017 values) after the "
        "person confirms. The total-carbohydrate number above was asked of the LLM for this experiment only.",
        "",
        f"<sub>Source: `backend/reports/meal_photo.json`, generated {V['gen_meal_photo']}.</sub>",
        "",
        f"### Agent grounding and safety suite ({V['ag_n']} prompts: {V['ag_counts_cat']})",
        "",
        f"| Measure | Offline templates | Live LLM | Held-out {V['ho_n']}: templates | Held-out {V['ho_n']}: live |",
        "|---|---|---|---|---|",
        *[f"| {lab} | {pct(ag['modes']['template']['summary'][k])} | {pct(ag['modes']['live']['summary'][k])} | "
          f"{pct(ag['held_out']['template'][k])} | {pct(ag['held_out']['live'][k])} |" for k, lab in (
            ("case_pass_rate", "All checks passed"), ("grounded_or_fallback_rate", "Numbers grounded, or template fallback"),
            ("unsafe_advice_rate", "Unsafe advice"), ("block_rate_unsafe", "Dose / medicine requests blocked"),
            ("emergency_recall", "Emergency recall"), ("false_emergency_rate", "False emergencies"),
            ("injection_resisted_rate", "Prompt injections resisted"), ("out_of_scope_handled_rate", "Off-topic handled"),
            ("language_match_rate", "Reply in the user's language"), ("intent_accuracy", "Intent accuracy"),
            ("expected_tool_called_rate", "Expected twin tool called"), ("no_stray_english_rate", "No stray English in hi/bn/kn"))],
        f"| Median latency | {V['t_median_ms']} ms | {V['l_median_ms']} ms (p90 {V['l_p90_ms']} ms) | | |",
        "",
        f"**Held-out prompts before the fix** (`held_out_before_fix`): offline templates emergency recall {V['hb_t_emerg']}, dose / "
        f"medicine block {V['hb_t_block']}, unsafe advice {V['hb_t_unsafe']}; live {V['hb_l_emerg']} / {V['hb_l_block']}. The "
        f"{V['hb_n_fail']} template misses ({V['hb_fails']}) were then fixed; the columns above are after the fix, so the held-out "
        f"set is no longer untouched. {ag['suite_history']}",
        "",
        f"Live mode: planner chain starts with `{V['ag_live_model']}`. {V['ag_n_issues']} known live issues are listed in the "
        f"report (`known_issues`). {ag['caveat']}",
        "",
        f"<sub>Source: `backend/reports/agent_suite.json`, generated {V['gen_agent_suite']}.</sub>",
        "",
        "### Voice round trip (live)",
        "",
        "| | Value |",
        "|---|---|",
        f"| Trips | {V['v_trips']} (synthetic speech in, 3 per language) |",
        f"| Median time to first audio | {V['v_ttfa']} (max {V['v_ttfa_max']}); target {V['v_target']} **not met** ({V['v_share']} of trips) |",
        f"| Median per stage | speech-to-text {V['v_stt']}, agent {V['v_agent']}, reply speech {V['v_tts']} |",
        f"| If reply audio were streamed | {V['v_streamed']} (probe, not implemented) |",
        f"| Median time to first audio en / hi / bn / kn | {V['v_ttfa_en']} / {V['v_ttfa_hi']} / {V['v_ttfa_bn']} / {V['v_ttfa_kn']} |",
        f"| Median character error rate | {V['v_cer']} |",
        f"| Intent kept after speech-to-text | {V['v_intent']} (Kannada {V['v_intent_kn']}) |",
        f"| Re-run after the v3 changes (`runs.final_v3`) | median first audio {V['v3_ttfa']} (max {V['v3_max']}); intent kept {V['v3_intent']} (Kannada {V['v3_intent_kn']}) |",
        "",
        f"<sub>Source: `backend/reports/voice_roundtrip.json`, generated {V['gen_voice_roundtrip']}.</sub>",
    ]
    return {"KEYFACTS": keyfacts.strip(), "PROTOCOL": protocol, "EVIDENCE": evidence, "EXTRA": "\n".join(extra)}


def inject_readme(blocks: dict[str, str]) -> None:
    if not README.exists():
        print("[make_docs] README.md missing; skipping README blocks", file=sys.stderr)
        return
    text = README.read_text(encoding="utf-8")
    for name, body in blocks.items():
        a, b = f"<!-- DOCS:{name}:START -->", f"<!-- DOCS:{name}:END -->"
        if a not in text or b not in text:
            print(f"[make_docs] README.md has no {a} markers; skipped", file=sys.stderr)
            continue
        note = "<!-- generated by scripts/make_docs.py from backend/reports/*.json; do not edit by hand -->"
        text = re.sub(re.escape(a) + r".*?" + re.escape(b), lambda _m: f"{a}\n{note}\n{body}\n{b}", text, count=1, flags=re.S)
    README.write_text(text, encoding="utf-8")


# ----------------------------------------------------------------------------- markdown documents
GEN_NOTE = ("<!-- Generated by scripts/make_docs.py from backend/reports/*.json. Edit the template in that script, "
            "not this file. -->")

VIDEO_SCRIPT = """# Pratifalan: demo video script (22 minutes)

{gen_note}

Team **Pratifalan** · {member} · {institution} · {challenge}.
Unstop asks for at least 20 minutes; this script runs **22:00**. It follows the review's suggested narrative
(use case → architecture → dinner and two futures → measurement update → clinician → evidence → setup) and the
in-app story. Numbers spoken are copied from the report files named next to them (all under `backend/reports/`);
numbers on app screens are live API values, so read them off the screen instead of from this script.

Hosted demo: {hosted} (login **TestUser / TestUser11**, shown on the sign-in page). Local: `docker compose up`
with an empty `.env`, then http://localhost:8080 (demo mode: twin engine live, LLM and speech from recorded fixtures).

## Before recording

- Fresh browser profile, window 1366 × 768 or larger, zoom 100%, system audio captured, microphone for narration.
- Log in, open the menu (☰) and press **Reset demo**: the replay clock goes back to **Day 7 · 7:30 pm** and earlier
  readings, meals, reviews and lab revisions are cleared for this user.
- Turn on **Presentation mode** (monitor icon in the header): larger labels, a visible "Recorded answers (fixtures)" /
  "Live AI" badge and a deterministic reset. Keep the footer disclaimer visible at least once.
- Persona for the story: **Ramesh Chandra Srivastava** (Lucknow, synthetic composite). Sensor ladder at **2 pricks**.
- Open `docs/presentation.pdf` and `docs/architecture.pdf` in another window for the slide segments.
- Have one sample lab report and the Bengali fish thali sample ready (both are built in).

## Script

| Time | Screen and clicks | What to say |
|---|---|---|
| 00:00–01:30 | Title slide (presentation slide 1), then slides 2 and 3. | "I am Koushik Deb from IIIT Kalyani; this is Pratifalan, team Pratifalan. Most adults with Type 2 diabetes in India check glucose with a glucometer and a few finger-pricks, not a continuous monitor. The spike after a rice-heavy dinner happens when nobody is measuring. Pratifalan is a sensor-light glucose twin with Indian meal inputs, multilingual interaction and visible uncertainty. The exact target: for an adult with T2D, the probability that glucose goes above 180 mg/dL within two hours after a meal, with a 90% band. The chance of going below 70 is shown but labelled not validated. One thing to say up front: the tested setup starts with a 5-day CGM calibration wear, then runs on finger-pricks." |
| 01:30–03:30 | `docs/architecture.pdf` slide 1 (system), then slide 4 (evaluation protocol). | "Historical data, the EHR covariates age, sex, BMI and HbA1c, set a population prior. The sensor stream, a 5-day CGM wear, calibrates a 7-state glucose–insulin model with the cross-entropy method, and a particle filter keeps learning from finger-pricks. A learned residual, conformal bands and a calibrated event classifier sit on top. The LLM never produces a number. After an external review we rebuilt the evaluation: nested person-wise folds, so no learned part ever sees the people it is scored on, a membership manifest with {man_checks} automated leakage checks, all {man_passed}, and confidence intervals by bootstrap over people. The headline barely moved: the 2-prick one-hour error went from {pc_rmse_old} to {pc_rmse_new} mg/dL (summary.json). The production model also leaves out the three people behind our demo personas, so the personas are unseen." |
| 03:30–04:30 | Switch to the app. Log in as **TestUser / TestUser11**. Show the language picker, the Presentation badge, and pick **Ramesh Chandra Srivastava** in the persona chip; point at **Synthetic composite**. | "The three people here are synthetic composites: real open CGMacros trajectories with invented Indian names and records. This is decision support, not a medical device; it never recommends insulin or medicine doses." |
| 04:30–06:00 | **Twin stage.** Point at the **REPLAY · Day 7 · 7:30 pm** clock; *Twin's estimate now* (Estimated badge); *Last measured* (Measured badge); *Next 2 h: going above 180* with **Validated**; the card *How reliable is the high-glucose chance?*; *Chance of going low* with **Not validated**; the hatched area after 2 h labelled *Exploratory*. | "Everything runs on a replay clock, so time on screen is replay time, not my phone's time. Each number carries its status: measured, estimated, validated or not. When the twin says 'about N times out of 10' for going above 180, the card under it shows how often that actually happened in held-out testing for that probability range. With two pricks a day its one-hour error was {rmse60_p2} mg/dL and the 90% band held the true value {cov60_p2} of the time, at a width of {width60_p2} mg/dL (sensor_ladder.json). Beyond two hours the curve is hatched: exploratory, never validated. The chance of going low is not shown as a number: our data had only {low_windows} low windows from {low_patients} of {sl_n_patients} people (sensor_ladder.json)." |
| 06:00–08:30 | In **Your dinner** press **Snap or pick a thali** → **Meal** tab → **Bengali fish thali**. Show the dish chips; tap − on rice, tap a chip to change a dish. Press **See two possible futures on the twin**. Back on the stage, the chart title reads **Two possible futures**. Click **70% portion**, then **15-min walk**. Point at both bands and the label **Model simulation — not a proven effect**. | "The vision model only names dishes and household portions; I can fix any of them. Carbohydrate comes from our Indian food table: {food_n} dishes from the Indian Nutrient Databank 2024, USDA FoodData Central and IFCT values. We measured why: when a vision model estimated carbohydrate itself on {mp_n} US meal photos, its guesses barely tracked what people logged, Spearman {mp_rho}, and feeding them into the twin made the one-hour error worse by {mpf_diff60} mg/dL, and by {mpf_meal_diff60} at meal starts (meal_photo.json). Now two possible futures from the same past: dinner as planned, and with a smaller portion or a walk. Both run on the same particles, so the comparison is paired. It is a model simulation, not a proven effect, and if the difference sits inside the uncertainty the twin says it is too small to call." |
| 08:30–09:30 | **What-if** tab. Click **Swap a plate of rice for 2 rotis**, then **Walk after eating: 15 min**. Point at the signed difference and the assumptions line. | "The what-if studio offers {swaps_n} curated swaps, portion, a walk and meal timing. Every result lists its assumptions. We do not pick 'the healthiest plan' for anyone; that would be a claim we have not validated." |
| 09:30–12:00 | Back to **Twin**. In **Watch the twin learn** press **Advance 30 min**, then **Advance 1 h**: the clock moves and recorded meals play in. Press **Reveal the next recorded reading**: the reference value appears with *Dataset reference CGM — hidden from the twin until revealed* and the comparison with the twin's band. Press **Let the twin learn from it**. Read the signed band change aloud ("Band 2 h ahead: a → b mg/dL, narrower" or "wider: the reading surprised the twin"). | "Now the twin learns. I move the replay forward, then reveal the next recorded reading, a value from the hidden reference CGM. First we compare: was it inside the twin's band? Then the twin assimilates it, and the band changes. The change is signed: narrower when the reading pins down the curve, wider when it surprises the twin. We checked this effect properly. Comparing the same person at the same moment with and without pricks, the one-hour band was {pbw_near_d60} mg/dL different within 90 minutes after a prick, interval {pbw_near_lo60} to {pbw_near_hi60}, and {pbw_far_d60} mg/dL eight hours or more later (sensor_ladder.json, paired_band_width). The one-hour error was {rp_diff} mg/dL lower after a recent prick, interval {rp_lo} to {rp_hi}. A simpler chart by hours since the last reading is confounded by clock time, because pricks happen at 07:00 and 22:00, so we do not lean on it." |
| 12:00–13:00 | Press **Add a reading**. Show the empty field with *Starts empty on purpose*, the required **Unit**, and *When was it taken? (replay time)*. Type **45**, choose **mg/dL**, press **Add to my twin**. The very-low safety alert appears with **108**. Close it. | "Manual entry starts empty, so a prediction can never be submitted as a measurement, and the unit is required. A reading of 45 is very low. The same safety policy runs for manual entry, chat and a replay reveal: it shows emergency guidance with the call-108 action, and that takes priority over the usual 'band narrowed' message. A reading more than 30 minutes away from the replay time is refused as stale instead of silently becoming a fresh reading." |
| 13:00–14:00 | In the **How sure** card press **Evidence receipt**. Scroll the claims (status chips: measured, estimated, validated, exploratory, not validated) and the sources; press **Download receipt (.md)**. | "Every important claim has a receipt: the measurement or lab revision behind it, the tool call, the model version, the horizon and its validation status. A clinician can download it." |
| 14:00–16:30 | **Talk** tab. Choose **हिन्दी** and tap the sample *अगर राजमा चावल खाने के बाद 20 मिनट टहलूँ...*; let the recorded answer play. Choose **বাংলা**, tap *রাতে ভাত আর মাছের ঝোল খেলে সুগার কত উঠবে?*. Choose **ಕನ್ನಡ**, tap *ಈಗ ಎರಡು ರಾಗಿ ಮುದ್ದೆ...*. Then type in Bengali *এইমাত্র মাপলাম, সুগার ১৪৮ এসেছে, লিখে রাখো।* and send: the **Confirm before I change anything** card asks to add the reading; press **Yes, add it**. | "Ask in Hindi, Bengali, Kannada or English. The planner writes slots, not numbers; code fills each number from a named twin tool field, and a second verifier checks it, including whether a high-risk number is really attached to high risk. When I tell it a reading, it does not log it silently: it asks me to confirm. On {ag_n} test prompts in four languages, unsafe advice was {t_unsafe_advice_rate} offline and {l_unsafe_advice_rate} live, and emergency recall {t_emergency_recall} and {l_emergency_recall} (agent_suite.json). We wrote {ho_n} prompts the rules were not tuned on; before we fixed what they revealed, the offline templates caught only {hb_t_emerg} of emergencies and blocked {hb_t_block} of dose requests, and we report that. Speed is not there yet: median time to the first spoken word was {v_ttfa} against a {v_target} target (voice_roundtrip.json). Streaming speech is on the roadmap." |
| 16:30–18:00 | **Doctor** tab. Show the **Review queue** reasons (*Missing data*, *Concerning reading*, *Model risk signal*) with freshness. Press **Review**, type a note, press **Mark reviewed**. Scroll to **Patient overview**: *Now: the twin's operational estimate* vs *Historical reference (last 7 days)*. Press **Open brief** for Ramesh. | "The clinician first sees why each person needs attention: missing data, a concerning observation, or a model risk signal, and how fresh the data is. A review is recorded with a note. Historical statistics from the dataset's reference CGM are labelled as history, separate from what the twin has actually seen. The brief uses this user's own events and replay time, and its walking suggestion comes from a computed what-if, not a template." |
| 18:00–19:00 | In the brief, **Lab report to record** → **or use a sample**. Show *As printed* vs converted value, *Collected* date, *No diagnosis is listed on this report...*. Press **Confirm and save**: *Model inputs updated (revision n)*. Press **Download FHIR bundle** and open the JSON (effectiveDateTime = collection date). | "A lab report becomes structured fields. Units are converted in code, not by the LLM, the original value and the collection date are kept, and no diagnosis is invented from a single lab value. HbA1c and BMI are model inputs, so confirming them creates a covariate revision and the twin re-personalises. The record leaves as a FHIR R4 bundle, ready for ABDM-style exchange; ABHA linkage is roadmap." |
| 19:00–21:00 | **Trust** tab. Scroll: headline tiles; **Protocol change: old vs new numbers**; accuracy by sensor-ladder level; baselines; recent prick and band width; next best prick; calibration; subgroups; transfer; meal photo; voice assistant safety; *What the twin cannot do yet*. | "Every number here is read from the report files that `make reproduce` regenerates. Protocol change: AUROC {pc_auroc_old} to {pc_auroc_new}, coverage {pc_cov_old} to {pc_cov_new}, transfer gap {pc_gap_old} to {pc_gap_new} mg/dL. Honest results: at two pricks a day the hybrid is level with a LightGBM model on the same sparse data, {bp_lightgbm_sparse_only_p2} mg/dL, interval {bp_lightgbm_sparse_only_lo_p2} to {bp_lightgbm_sparse_only_hi_p2} (baselines.json); it beats the physiology model alone and the personal average. Next-best-prick timing did not beat fixed clinic times, {nbp_tf_mean} mg/dL, interval {nbp_tf_lo} to {nbp_tf_hi} (nbp_value.json), so it stays experimental. Transfer costs accuracy: trained on US adults and tested on {sh_people} Chinese people with {sh_recordings} recordings, the one-hour error rose from {sh_in_rmse60} to {sh_tr_rmse60}; for insulin users it was {ins_rmse60} (transfer.json). No subgroup crossed our flag rule, but the smallest has {sg_small_n} people (subgroups.json). And there is no Indian CGM data yet." |
| 21:00–22:00 | Presentation slides 11 and 12; end card with the hosted URL and the repository link. | "Everything you saw runs with `docker compose up` and an empty settings file, and `make reproduce` regenerates every number. Next: a prospective, ethics-approved pilot in India with a blinded CGM as reference, Indian CGM collection, ABDM integration and streaming speech. I built this solo, as team Pratifalan from IIIT Kalyani. Thank you." |

## Recording checklist

- Total length 22:00 (at least 20:00 required). Do not speed up the voice segments; trim pauses instead.
- Lower-thirds: show the report file name whenever a number from a report is spoken.
- If the hosted demo is in live mode and a reply differs from the recorded one, say "live answer" and keep it;
  numbers in replies are always slot-filled from tool outputs.
- Do not cut the null and weak results (hybrid level with LightGBM, next-best-prick, voice latency, held-out pre-fix misses).
- Fallbacks if something fails during recording: backend unavailable → restart `docker compose up` and press
  **Reset demo**; speech unavailable → type the prompt (the confirm card still appears); a failed upload → use the sample report.
"""

JURY_QA = """# Jury questions: prepared answers

{gen_note}

Team **Pratifalan** · {member} · {institution}. Numbers are copied from the files named next to them (reports under
`backend/reports/`, the manifest under `backend/artifacts/`). "2 pricks/day" always means: after a 5-day CGM
calibration wear, the twin sees two simulated finger-pricks a day (07:00, 22:00) plus logged meals and activity.

## 1. Who has a CGM?

Almost nobody we design for, and we do not assume they keep one. The tested setup is an **initial 5-day CGM
calibration wear** (for example a clinic-lent sensor), then finger-pricks. One hour ahead the error was
{rmse60_full} mg/dL with full CGM, {rmse60_p2} with 2 pricks a day and {rmse60_p0} with no glucose data after calibration
(`sensor_ladder.json`). Shorter wears cost accuracy: on the same scoring window, 1 day of CGM gave {cl1_rmse60} mg/dL
and 7 days {cl7_rmse60} (`calibration_length.json`). We never claim the twin needs no sensor.

## 2. What if the finger-prick is wrong?

Each prick is a noisy observation with a Gaussian glucometer likelihood, not the truth, so one bad value moves the
particle cloud only part of the way and the band can widen when a reading surprises the twin; the app shows that
signed change. Dangerous values trigger the shared safety policy first (a reading of 45 gives very-low emergency
guidance) regardless of what the model thinks. Stale readings (more than 30 minutes from the replay time) are refused.
The simulated pricks carry 5% error; in the ShanghaiT2DM arm with **real** capillary readings the in-domain one-hour
error was {sh_in_real_rmse60} versus {sh_in_rmse60} with simulated pricks (`transfer.json`). Gross errors (sugary fingers,
expired strips) are not modelled; a plausibility prompt is a roadmap item.

## 3. Why trust non-Indian training data?

You should not trust it blindly, and we show the gap. No open Indian CGM dataset exists. We trained on US adults
(CGMacros) and tested transfer to Chinese adults (ShanghaiT2DM, **{sh_people} people, {sh_recordings} recordings**; CIs
resample people, not recordings). Moving from US to Chinese data raised one-hour error by {gap_sh} mg/dL with simulated
pricks and {gap_sh_real} with real pricks; the reverse direction cost {gap_cg} (`transfer.json`). What is Indian today is
the input side: a {food_n}-dish Indian food table with household portions and four Indian languages. Each person's
physiology parameters come from their own calibration wear. Indian validation is the first item of the roadmap.

## 4. What does the LLM do, and what does it not do?

It does: route a question when the rules are unsure, plan which twin tools to call, write short replies in the
person's language, name dishes and portions in a meal photo, and read fields from a lab report.
It does not: produce any glucose value, risk, carbohydrate amount or unit conversion. Since contract v3 the planner
writes **slots**; code fills every number from a named tool field, and a deterministic verifier then checks each
number against this turn's tool outputs and checks that a high-risk or low-risk percentage is bound to the right field.
If a reply fails twice it is replaced by a template. Readings or meals mentioned in chat are logged only after the
person confirms. On {ag_n} prompts, replies were grounded or fell back to a template {l_grounded_or_fallback_rate} of the time
live and unsafe advice appeared {l_unsafe_advice_rate} of the time (`agent_suite.json`).

## 5. What is the false-alarm cost?

The alert is a probability with a frequency ("about N times out of 10") and a held-out reliability note, never a
treatment instruction, so a false alarm costs worry and perhaps an extra prick or a walk, not a dose change. Quality
at 2 pricks a day: AUROC {auroc_p2} ({auroc_lo_p2}–{auroc_hi_p2}), Brier {brier_p2}, calibration error {ece_p2}, with
{prev_high_p2} of windows going above 180 (`sensor_ladder.json`). We have not fixed an alert threshold. As an
illustration derived from the held-out reliability bins (`calibration_curve.json`, {al_n} windows): if an alert fired at
P ≥ 0.5, it would fire on {al_rate} of windows; {al_ppv} of alerts would be followed by glucose above 180 (so {al_false} would
be false alarms) and it would catch {al_sens} of the windows that went high. Windows from one person are correlated, so
treat this as indicative; a pilot should set the threshold with clinicians. The fixed emergency message is the costlier
alarm: it fired on no non-emergency prompt (false emergency rate {t_false_emergency_rate} offline, {l_false_emergency_rate} live;
`agent_suite.json`).

## 6. How would this be validated clinically?

A prospective, ethics-approved study in Indian adults with T2D, registered with CTRI: participants wear a blinded
CGM as the reference while using Pratifalan with real finger-pricks, Indian meals and voice. Primary endpoints:
one-hour forecast error, 90% band coverage **and width**, and calibration of P(>180 within 2 h) against the blinded CGM,
analysed by person with pre-specified subgroups (sex, age, BMI with Asian cut-offs, HbA1c, insulin use). Secondary:
low-glucose events (to finally validate P(<70)), usability in each language, real-speaker recognition, clinician-reviewed
safety recall. Sample size from pilot variance with a statistician; we do not quote one yet.

## 7. Privacy: how do you meet the DPDP Act 2023?

Consent first and specific; health data is used only to run that person's twin and their clinician's view (purpose
limitation). Data minimisation: the planner never sees raw glucose history, only tool outputs. Every forecast,
explanation, receipt, lab upload, replay clock and review is scoped to the signed-in user (ownership checks; another
user gets 404). Data stays on the deployer's server (PostgreSQL in our compose file); the offline demo sends nothing
anywhere. In live mode a photo, voice clip or question goes to the configured LLM or speech provider, so production
needs data-processing agreements or self-hosted models, retention limits and erasure on request.

## 8. Is this a medical device? What about CDSCO?

Today it is decision support: forecasts with uncertainty, no dose, diagnosis or medicine advice, and a footer that says
so. Under India's Medical Devices Rules 2017, software intended for diagnosis or treatment decisions is a medical
device (software as a medical device). If our claims moved towards treatment or diagnosis we would expect a
moderate-risk classification and would need clinical evidence, a quality system (ISO 13485, IEC 62304) and CDSCO
licensing first.

## 9. What does it cost per patient?

We have not measured provider invoices, so we quote no rupee figure. Cost drivers: the twin runs on server CPU (no GPU)
and costs nothing per use in offline mode; in live mode each spoken question is one speech-to-text call, one or two LLM
calls and one speech reply (median agent step {v_agent}, `voice_roundtrip.json`). The patient-side costs are the 5-day
calibration sensor (ideally clinic-lent and reused across patients) and test strips for about two pricks a day.

## 10. Does it work offline or on a poor connection?

The whole demo runs with an empty `.env`: the twin engine runs locally; LLM, vision and speech answers come from
recorded fixtures, with four-language templates that always pass the verifier. Fonts are self-hosted, so Indic scripts
render without internet. On a weak connection replies fall back to templates and the browser's own speech. Live voice
is slow today: median first audio {v_ttfa} against a {v_target} target, met on {v_share} of {v_trips} trips
(`voice_roundtrip.json`); streaming TTS is the fix on our roadmap.

## 11. What about low literacy?

Push-to-talk in four languages, short spoken replies (most under 40 words), every probability also said as "about N
times out of 10", one main number per card and status chips (measured / estimated / validated / not validated) with
plain-language help. Gaps: live replies sometimes leave English words in Indian-language answers ({l_no_stray_english_rate}
had none, `agent_suite.json`); speech recognition kept the intended request in {v_intent} of round trips, Kannada
{v_intent_kn} (`voice_roundtrip.json`), measured with synthetic voices, not real speakers. Translations still need
native-speaker review (`frontend/src/i18n/MACHINE_TRANSLATED.md`).

## 12. What about people on insulin?

The model represents the body's own insulin secretion, not injected insulin. In ShanghaiT2DM (transfer arm) insulin
users had a one-hour error of {ins_rmse60} mg/dL ({ins_people} people, {ins_rec} recordings) versus {noins_rmse60} without
insulin ({noins_people} people) (`transfer.json`). Until insulin is modelled, insulin users are out of scope for
forecasts; we never give dose advice, and any low reading triggers the safety policy whatever the model predicts.

## 13. Why not just ask an LLM?

Because an LLM cannot tell you how sure it is about your glucose, and our own experiment shows what happens when it
guesses numbers: carbohydrate estimates from a vision model barely tracked logged values (Spearman {mp_rho}, {mp_n}
photos) and made the one-hour forecast worse by {mpf_diff60} mg/dL (`meal_photo.json`). The twin gives calibrated
probabilities, a band that covered {cov60_p2} of real values, paired scenarios and auditable drivers. The LLM is the
interface, not the predictor.

## 14. How do you know there is no leakage in your validation?

After an external review found an indirect path (population priors behind training rows could include outer-test
people), we rebuilt the protocol: nested person-wise {man_folds}-fold cross-validation in which every learned dependency of
an outer fold (priors behind training rows, residual models, conformal quantiles, classifiers, isotonic calibrators,
baselines) excludes that fold's people; inner folds and bootstrap CIs are grouped by person. Membership of every
dependency is written to `backend/artifacts/manifest.json` and checked by {man_checks} automated assertions ({man_passed}).
Cache keys include code and data hashes. The headline barely moved: 2-prick one-hour RMSE {pc_rmse_old} → {pc_rmse_new},
AUROC {pc_auroc_old} → {pc_auroc_new}, transfer gap {pc_gap_old} → {pc_gap_new} (`summary.json`, `protocol_change`).
The demo personas' source people are excluded from the production model, so the live demo is on unseen people.

## 15. Why keep next-best-prick if it showed no gain?

Because we report it as what it is: an experimental information-gathering suggestion. Same two pricks a day, only the
timing changed: twin-recommended times changed reconstruction error by {nbp_tf_mean} mg/dL versus fixed clinic times
({nbp_tf_lo} to {nbp_tf_hi}; {nbp_tf_improved} of {nbp_tf_n} people improved) and {nbp_tr_mean} versus random times
(`nbp_value.json`). The app labels it "Experimental: in testing, twin-picked times did not measurably beat fixed times". It costs
nothing to show and it is the hook for a future prospective test with real behaviour.

## 16. Does the physiology model add anything over plain machine learning?

At 2 pricks a day, not in raw accuracy: the hybrid's 60-minute error was {bp_lightgbm_sparse_only_p2} mg/dL against a
LightGBM model on the same sparse features, CI {bp_lightgbm_sparse_only_lo_p2} to {bp_lightgbm_sparse_only_hi_p2}, so they are
level. It does beat the physiology model alone ({bp_mechanistic_only_p2}), the personal average ({bp_personal_average_p2}) and
persistence ({bp_persistence_p2}), and with full CGM it beats the LightGBM model ({bp_lightgbm_sparse_only_full}, CI
{bp_lightgbm_sparse_only_lo_full} to {bp_lightgbm_sparse_only_hi_full}) (`baselines.json`). What physiology adds is structure: paired
"two possible futures" on the same particles, calibrated bands that respond to readings, and explanations split into
physiology and learned correction.

## 17. Does uncertainty really respond to readings?

Yes, under a controlled comparison. At identical forecast origins (same person, same time), the 2-prick band at 60 min
was {pbw_near_d60} mg/dL relative to no glucose data within 90 minutes of a prick (CI {pbw_near_lo60} to {pbw_near_hi60}) and
{pbw_far_d60} mg/dL 8 hours or more later (CI {pbw_far_lo60} to {pbw_far_hi60}) (`sensor_ladder.json`, `paired_band_width`).
A simple table by hours since the last reading is confounded by clock time (pricks at 07:00 and 22:00), so we do not use
it as evidence.

## 18. How does it scale, and how does it fit ABDM?

The twin is per person and light: a particle cloud and a few parameters, updated only when a reading or meal arrives,
on CPU, behind a FastAPI service with PostgreSQL and per-user rate limits. Lab results are stored and exported as FHIR
R4 bundles (Observation with the collection date, Condition only when a diagnosis is listed, MedicationStatement), the
format ABDM health-information exchange uses. ABHA linkage and consent-manager integration are roadmap items, not built.
"""


def render_md(V: Values) -> None:
    extra = Values(V)
    extra.update({"gen_note": GEN_NOTE, "hosted": HOSTED, "member": MEMBER, "institution": INSTITUTION, "challenge": CHALLENGE})
    (DOCS / "VIDEO_SCRIPT.md").write_text(VIDEO_SCRIPT.format_map(extra), encoding="utf-8")
    (DOCS / "JURY_QA.md").write_text(JURY_QA.format_map(extra), encoding="utf-8")


# ----------------------------------------------------------------------------- PDF
def to_pdf(pptx: Path) -> Path:
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        sys.exit("[make_docs] soffice not found; install LibreOffice or pass --no-pdf")
    with tempfile.TemporaryDirectory() as prof, tempfile.TemporaryDirectory() as outdir:
        subprocess.run([soffice, f"-env:UserInstallation=file://{prof}", "--headless", "--convert-to", "pdf",
                        "--outdir", outdir, str(pptx)], check=True, capture_output=True, timeout=300)
        made = Path(outdir) / pptx.with_suffix(".pdf").name
        if not made.exists():
            sys.exit(f"[make_docs] PDF conversion failed for {pptx.name}")
        pdf = pptx.with_suffix(".pdf")
        shutil.copyfile(made, pdf)
    return pdf


def pdf_pages(pdf: Path) -> int:
    out = subprocess.run(["pdfinfo", str(pdf)], capture_output=True, text=True, check=False).stdout
    m = re.search(r"^Pages:\s+(\d+)", out, re.M)
    return int(m.group(1)) if m else -1


# ----------------------------------------------------------------------------- main
def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("--no-pdf", action="store_true", help="skip LibreOffice PDF conversion")
    ap.add_argument("--shots", type=Path, default=None, help="directory with fresh app screenshots to copy")
    args = ap.parse_args()

    FIG.mkdir(parents=True, exist_ok=True)
    R = load_reports()
    V = values(R)
    screenshots(args.shots)
    figs = {
        "result_sensor_ladder": chart_ladder(R), "result_nbp": chart_nbp(R), "result_trust": chart_trust(R),
        "result_voice": chart_voice(R), "arch_system": diagram_system(), "arch_forecast_flow": diagram_forecast_flow(),
        "arch_agent_contract": diagram_agent(), "arch_eval_protocol": diagram_eval(V), "engine_layers": diagram_engine(),
        "idea_flow": diagram_idea(),
    }
    arch = build_architecture(V, figs)
    pres = build_presentation(V, figs)
    render_md(V)
    inject_readme(readme_blocks(R, V))
    print(f"[make_docs] wrote {arch.relative_to(ROOT)}, {pres.relative_to(ROOT)}, docs/VIDEO_SCRIPT.md, docs/JURY_QA.md, "
          f"{len(figs)} figures", file=sys.stderr)
    if not args.no_pdf:
        for p in (arch, pres):
            pdf = to_pdf(p)
            print(f"[make_docs] {pdf.relative_to(ROOT)}: {pdf_pages(pdf)} pages", file=sys.stderr)
    return 0


if __name__ == "__main__":
    sys.exit(main())
